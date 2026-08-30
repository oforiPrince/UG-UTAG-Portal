from __future__ import annotations

import csv
import io
import re
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from uuid import UUID
from zipfile import BadZipFile

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from utag_api.database import new_id
from utag_api.errors import ApiError
from utag_api.models import (
    Announcement,
    Conversation,
    ConversationInvite,
    ConversationMember,
    Document,
    Message,
    OrganizationUnit,
    Role,
    User,
    UserRole,
)
from utag_api.security import hash_password, normalize_email
from utag_api.seed_data import (
    _referenced_unit_ids,
    _unit_depth,
    organization_name_key,
    seed_organization,
)
from utag_api.services.member_import import (
    clean_cell,
    normalized_header,
    normalized_row,
    verify_xlsx_archive,
)
from utag_api.services.system_chat_groups import sync_system_chat_groups

MAX_ROSTER_BYTES = 20_000_000
MAX_ROSTER_ROWS = 20_000

HEADER_SIGNALS = frozenset(
    {
        "college",
        "department",
        "email",
        "full_name",
        "other_name",
        "school",
        "staff_id",
        "surname",
    }
)

HONORIFIC_PREFIX = re.compile(
    r"^(?:prof(?:essor)?|dr|mr|mrs|ms|miss|rev|eng(?:r)?|sir|hon)\.?\s+",
    re.IGNORECASE,
)

_SMALL_WORDS = frozenset({"of", "and", "for", "the"})
_SCHOOL_MARKERS = (
    "administration",
    "campus",
    "centre",
    "center",
    "chancellor",
    "college",
    "directorate",
    "faculty",
    "institute",
    "librarian",
    "library",
    "office",
    "school",
    "secretariat",
)
_NON_DEPARTMENT_MARKERS = (*_SCHOOL_MARKERS, "research")
_SECRETARIAT_SUFFIX = re.compile(r"\s+secretariat$", re.IGNORECASE)
_UNIT_DISPLAY_PAIRS: tuple[tuple[str, str], ...] = (
    ("Centre for Urban Mgt Studies", "Centre for Urban Management Studies"),
    ("College of Basic & Applied Sc.", "College of Basic and Applied Sciences"),
    ("Dept. of African and Asian Language", "Department of African and Asian Languages"),
    (
        "Human Resource and Oganisation Development Directorate",
        "Human Resource and Organisation Development Directorate",
    ),
    (
        "Inst. of Stat. Soc.& Econ. Res",
        "Institute of Statistical, Social and Economic Research",
    ),
    (
        "Legon Centre for Int. Affairs",
        "Legon Centre for International Affairs and Diplomacy",
    ),
    (
        "Noguchi Mem. Inst. for Med.Res",
        "Noguchi Memorial Institute for Medical Research",
    ),
    ("Ofice of the Universtiy Librarian", "Office of the University Librarian"),
    (
        "Regional Inst. of Pop. Studies",
        "Regional Institute for Population Studies",
    ),
    ("Sch of Biological Sciences", "School of Biological Sciences"),
    ("Sch of Educ & Leadership", "School of Education and Leadership"),
    (
        "Sch. of Biomed & Allied Health",
        "School of Biomedical and Allied Health Sciences",
    ),
    (
        "Sch. of Info. & Comm. Studies",
        "School of Information and Communication Studies",
    ),
    (
        "Sch. of Physical & Math. Sc.",
        "School of Physical and Mathematical Sciences",
    ),
    ("School of Cont. & Dist. Educ.", "School of Continuing and Distance Education"),
    ("School of Engineering", "School of Engineering Sciences"),
    ("University of Ghana Business School", "University of Ghana Business School"),
    ("University of Ghana Dental School", "University of Ghana Dental School"),
    ("University of Ghana Medical School", "University of Ghana Medical School"),
    ("GEMP Office", "GEMP Office"),
    (
        "Department of Operations and MIS",
        "Department of Operations and Management Information Systems",
    ),
    (
        "Department of Organisation and HR Management",
        "Department of Organisation and Human Resource Management",
    ),
)

TITLE_ALIASES = {
    "dame": "Dame",
    "doctor": "Dr.",
    "dr": "Dr.",
    "dralhaji": "Dr. (Alhaji)",
    "drmiss": "Dr. (Miss)",
    "drmrs": "Dr. (Mrs.)",
    "eng": "Eng.",
    "engr": "Eng.",
    "hon": "Hon.",
    "miss": "Miss",
    "mister": "Mr.",
    "mr": "Mr.",
    "mrs": "Mrs.",
    "ms": "Ms.",
    "mx": "Mx.",
    "prof": "Prof.",
    "professor": "Prof.",
    "professormrs": "Prof. (Mrs.)",
    "profmrs": "Prof. (Mrs.)",
    "rev": "Rev.",
    "reverend": "Rev.",
    "sir": "Sir",
}

RANK_ALIASES = {
    "assistantlecturer": "Assistant Lecturer",
    "assistantresearchfellow": "Assistant Research Fellow",
    "associateprof": "Associate Professor",
    "associateprofessor": "Associate Professor",
    "assocprofessor": "Associate Professor",
    "asstlecturer": "Assistant Lecturer",
    "assistresearchfellow": "Assistant Research Fellow",
    "dean": "Dean",
    "director": "Director",
    "lecturer": "Lecturer",
    "librarian": "Librarian",
    "principalresearchfellow": "Principal Research Fellow",
    "professor": "Professor",
    "provicechancellor": "Pro-Vice-Chancellor",
    "researchassociate": "Research Associate",
    "researchfellow": "Research Fellow",
    "seniorlecturer": "Senior Lecturer",
    "seniorlibrarian": "Senior Librarian",
    "seniorresearchfellow": "Senior Research Fellow",
    "seniorlcturer": "Senior Lecturer",
    "snrlecturer": "Senior Lecturer",
    "snrresearchfellow": "Senior Research Fellow",
    "tutor": "Tutor",
    "vicechancellor": "Vice-Chancellor",
    "visitingscholar": "Visiting Scholar",
}

GENDER_ALIASES = {
    "f": "Female",
    "female": "Female",
    "m": "Male",
    "male": "Male",
}


@dataclass(frozen=True, slots=True)
class RosterRow:
    row_number: int
    sheet: str
    staff_id: str
    email: str
    title: str
    other_name: str
    surname: str
    gender: str
    academic_rank: str
    phone_number: str
    college: str
    school: str
    department: str


@dataclass(frozen=True, slots=True)
class MemberRosterSyncResult:
    rows_total: int
    units_created: int
    units_updated: int
    members_relinked: int
    titles_filled: int
    ranks_filled: int
    staff_ids_filled: int
    members_created: int
    members_unchanged: int
    unmatched_existing_members: int
    units_pruned: int
    units_retained: int
    chat_groups_created: int
    chat_memberships_added: int
    chat_memberships_removed: int
    conflicts: tuple[str, ...]
    row_issues: tuple[str, ...]
    retained_units: tuple[str, ...]
    affected_user_ids: tuple[UUID, ...]


def compact_token(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", value.casefold())


def _unit_display_names() -> dict[str, str]:
    mapping: dict[str, str] = {}
    for variant, display in _UNIT_DISPLAY_PAIRS:
        mapping[organization_name_key(variant)] = display
        mapping[organization_name_key(display)] = display
    return mapping


UNIT_DISPLAY_NAMES = _unit_display_names()


def same_unit(left: str, right: str) -> bool:
    return bool(left and right) and organization_name_key(left) == organization_name_key(right)


def strip_secretariat(name: str) -> str:
    stripped = _SECRETARIAT_SUFFIX.sub("", name).strip()
    return stripped or name


def display_unit_name(value: str) -> str:
    cleaned = " ".join(value.split())
    if not cleaned:
        return ""
    letters = [character for character in cleaned.replace("&", "") if character.isalpha()]
    if letters and (
        all(character.isupper() for character in letters)
        or all(character.islower() for character in letters)
    ):
        words = cleaned.replace("&", " and ").casefold().split()
        titled: list[str] = []
        for index, word in enumerate(words):
            if index > 0 and word in _SMALL_WORDS:
                titled.append(word)
            else:
                titled.append(_title_unit_token(word))
        return " ".join(titled)
    return " ".join(cleaned.replace("&", " and ").split())


def _title_unit_token(word: str) -> str:
    if word.startswith("(") and word.endswith(")") and len(word) > 2:
        inner = word[1:-1]
        if inner.isalpha() and 2 <= len(inner) <= 6:
            return f"({inner.upper()})"
        return f"({_title_hyphen_slash(inner)})"
    return _title_hyphen_slash(word)


def _title_hyphen_slash(word: str) -> str:
    return "/".join(
        "-".join((part[:1].upper() + part[1:]) if part else part for part in chunk.split("-"))
        for chunk in word.split("/")
    )


def canonical_unit_name(value: str, unit_type: str) -> str:
    cleaned = display_unit_name(value)
    if not cleaned:
        return ""
    aliased = UNIT_DISPLAY_NAMES.get(organization_name_key(cleaned))
    if aliased:
        return aliased
    key = organization_name_key(cleaned)
    if unit_type == "college":
        if "college" in key or "administration" in key:
            return cleaned
        return f"College of {cleaned}"
    if unit_type == "school":
        if any(marker in key for marker in _SCHOOL_MARKERS):
            return cleaned
        return f"School of {cleaned}"
    if unit_type == "department":
        if "department" in key or any(marker in key for marker in _NON_DEPARTMENT_MARKERS):
            return cleaned
        return f"Department of {cleaned}"
    return cleaned


def normalize_affiliation(
    college: str,
    school: str,
    department: str,
) -> tuple[str, str, str]:
    college_name = canonical_unit_name(college, "college")
    school_name = canonical_unit_name(school, "school")
    department_name = canonical_unit_name(department, "department")
    if college_name and (not school_name or same_unit(college_name, school_name)):
        lifted = strip_secretariat(department_name)
        school_name = canonical_unit_name(lifted, "school") if lifted else ""
        department_name = ""
    elif (
        department_name
        and school_name
        and same_unit(strip_secretariat(department_name), school_name)
    ):
        department_name = ""
    elif department_name:
        stripped = strip_secretariat(department_name)
        if stripped != department_name:
            department_name = canonical_unit_name(stripped, "department")
    return college_name, school_name, department_name


def canonicalize_title(value: str) -> str:
    stripped = value.strip()
    if not stripped:
        return ""
    mapped = TITLE_ALIASES.get(compact_token(stripped))
    return (mapped or stripped)[:30]


def canonicalize_rank(value: str) -> str:
    stripped = value.strip()
    if not stripped:
        return ""
    mapped = RANK_ALIASES.get(compact_token(stripped))
    return (mapped or stripped)[:120]


def canonicalize_gender(value: str) -> str:
    return GENDER_ALIASES.get(compact_token(value), value.strip())[:30]


def split_person_name(
    *,
    full_name: str,
    other_name: str,
    surname: str,
) -> tuple[str, str]:
    if other_name.strip() and surname.strip():
        return other_name.strip()[:120], surname.strip()[:120]
    source = HONORIFIC_PREFIX.sub("", (full_name or f"{other_name} {surname}").strip())
    parts = [part for part in source.split() if part]
    if not parts:
        return other_name.strip()[:120], surname.strip()[:120]
    if len(parts) == 1:
        token = parts[0][:120]
        if surname.strip():
            return token, surname.strip()[:120]
        if other_name.strip():
            return other_name.strip()[:120], token
        return token, token
    return " ".join(parts[:-1])[:120], parts[-1][:120]


def _is_header_row(headers: Sequence[str]) -> bool:
    fields = set(headers)
    if "email" not in fields:
        return False
    return len(fields.intersection(HEADER_SIGNALS)) >= 2


def _cell(row: Mapping[str, object], key: str) -> str:
    value = row.get(key, "")
    if isinstance(value, list):
        return " ".join(str(item) for item in value).strip()
    return str(value or "").strip()


def _roster_row_from_values(
    *,
    row_number: int,
    sheet: str,
    headers: list[str],
    values: Sequence[object],
) -> RosterRow | None:
    if not any(clean_cell(value) for value in values):
        return None
    raw = normalized_row(headers, values)
    other_name, surname = split_person_name(
        full_name=_cell(raw, "full_name"),
        other_name=_cell(raw, "other_name"),
        surname=_cell(raw, "surname"),
    )
    email = _cell(raw, "email")
    staff_id = _cell(raw, "staff_id")
    college, school, department = normalize_affiliation(
        _cell(raw, "college"),
        _cell(raw, "school"),
        _cell(raw, "department"),
    )
    if not any((email, staff_id, other_name, surname, college, school, department)):
        return None
    return RosterRow(
        row_number=row_number,
        sheet=sheet,
        staff_id=staff_id[:40],
        email=email,
        title=canonicalize_title(_cell(raw, "title")),
        other_name=other_name,
        surname=surname,
        gender=canonicalize_gender(_cell(raw, "gender")),
        academic_rank=canonicalize_rank(_cell(raw, "academic_rank")),
        phone_number=_cell(raw, "phone_number")[:40],
        college=college,
        school=school,
        department=department,
    )


def _rows_from_table(
    *,
    sheet: str,
    raw_rows: Sequence[Sequence[object]],
) -> list[RosterRow]:
    header_index: int | None = None
    headers: list[str] = []
    for index, row in enumerate(raw_rows):
        candidate = [normalized_header(value) for value in row]
        if _is_header_row(candidate):
            header_index = index
            headers = candidate
            break
    if header_index is None:
        return []
    parsed: list[RosterRow] = []
    for offset, values in enumerate(raw_rows[header_index + 1 :]):
        item = _roster_row_from_values(
            row_number=header_index + 2 + offset,
            sheet=sheet,
            headers=headers,
            values=values,
        )
        if item is not None:
            parsed.append(item)
    return parsed


def parse_member_roster_bytes(filename: str, data: bytes) -> list[RosterRow]:
    if not data:
        raise ApiError(422, "import_empty", "The roster file is empty")
    if len(data) > MAX_ROSTER_BYTES:
        raise ApiError(413, "import_too_large", "Roster files must be 20 MB or smaller")
    suffix = Path(filename).suffix.casefold()
    if suffix == ".csv":
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ApiError(
                422, "import_encoding_invalid", "CSV files must use UTF-8 encoding"
            ) from exc
        reader = csv.reader(io.StringIO(text))
        rows = _rows_from_table(sheet="Sheet1", raw_rows=list(reader))
    elif suffix == ".xlsx":
        verify_xlsx_archive(data)
        try:
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except (OSError, ValueError, KeyError, BadZipFile) as exc:
            raise ApiError(422, "workbook_invalid", "The Excel workbook could not be read") from exc
        try:
            rows = [
                item
                for worksheet in workbook.worksheets
                for item in _rows_from_table(
                    sheet=worksheet.title,
                    raw_rows=list(worksheet.iter_rows(values_only=True)),
                )
            ]
        finally:
            workbook.close()
    else:
        raise ApiError(415, "import_type_invalid", "Choose a CSV or XLSX roster file")
    if not rows:
        raise ApiError(422, "import_empty", "The roster file has no member rows")
    if len(rows) > MAX_ROSTER_ROWS:
        raise ApiError(413, "import_row_limit", f"Rosters are limited to {MAX_ROSTER_ROWS} rows")
    return rows


def parse_member_roster(path: Path) -> list[RosterRow]:
    return parse_member_roster_bytes(path.name, path.read_bytes())


def organization_structure_from_rows(
    rows: Sequence[RosterRow],
) -> dict[str, dict[str, tuple[str, ...]]]:
    tree: dict[str, dict[str, list[str]]] = {}
    college_by_key: dict[str, str] = {}
    school_by_key: dict[tuple[str, str], str] = {}
    for row in rows:
        if not row.college or not row.school:
            continue
        college_key = organization_name_key(row.college)
        college_name = college_by_key.setdefault(college_key, row.college)
        schools = tree.setdefault(college_name, {})
        school_key = organization_name_key(row.school)
        school_name = school_by_key.setdefault((college_key, school_key), row.school)
        departments = schools.setdefault(school_name, [])
        if not row.department:
            continue
        department_key = organization_name_key(row.department)
        if department_key not in {organization_name_key(name) for name in departments}:
            departments.append(row.department)
    return {
        college: {
            school: tuple(sorted(departments, key=str.casefold))
            for school, departments in sorted(schools.items(), key=lambda item: item[0].casefold())
        }
        for college, schools in sorted(tree.items(), key=lambda item: item[0].casefold())
    }


def render_organization_structure(
    structure: dict[str, dict[str, tuple[str, ...]]],
) -> str:
    lines = ["{"]
    colleges = list(structure.items())
    for college_index, (college, schools) in enumerate(colleges):
        college_comma = "," if college_index < len(colleges) - 1 else ""
        lines.append(f'    "{college}": {{')
        school_items = list(schools.items())
        for school_index, (school, departments) in enumerate(school_items):
            school_comma = "," if school_index < len(school_items) - 1 else ""
            if not departments:
                lines.append(f'        "{school}": (){school_comma}')
                continue
            lines.append(f'        "{school}": (')
            for department in departments:
                lines.append(f'            "{department}",')
            lines.append(f"        ){school_comma}")
        lines.append(f"    }}{college_comma}")
    lines.append("}")
    return "\n".join(lines)


def _find_unit(
    units: Sequence[OrganizationUnit],
    *,
    unit_type: str,
    name: str,
    parent_id: UUID | None,
) -> OrganizationUnit | None:
    name_key = organization_name_key(name)
    matches = [
        item
        for item in units
        if item.unit_type == unit_type and organization_name_key(item.name) == name_key
    ]
    under_parent = [item for item in matches if item.parent_id == parent_id]
    if len(under_parent) == 1:
        return under_parent[0]
    if len(matches) == 1:
        return matches[0]
    return None


def _blank(value: str | None) -> bool:
    return not (value or "").strip()


def _valid_email(value: str) -> bool:
    return "@" in value and " " not in value.strip()


def _row_label(row: RosterRow) -> str:
    location = f"row {row.row_number}"
    if row.sheet and row.sheet != "Sheet1":
        return f"{row.sheet} {location}"
    return location


async def _prune_leftover_units(
    db: AsyncSession,
    *,
    canonical_ids: frozenset[UUID],
    users: Sequence[User],
) -> tuple[int, tuple[str, ...]]:
    units = list((await db.scalars(select(OrganizationUnit))).all())
    units_by_id = {item.id: item for item in units}
    leftover_ids = {
        item.id
        for item in units
        if item.unit_type in {"college", "school", "department"} and item.id not in canonical_ids
    }
    announcements = list((await db.scalars(select(Announcement))).all())
    documents = list((await db.scalars(select(Document))).all())
    system_groups = list(
        (
            await db.scalars(select(Conversation).where(Conversation.direct_key.like("system:%")))
        ).all()
    )
    system_group_by_unit_id: dict[UUID, Conversation] = {}
    for conversation in system_groups:
        direct_key = conversation.direct_key or ""
        prefix, _, raw_unit_id = direct_key.rpartition(":")
        if prefix not in {"system:school", "system:department"}:
            continue
        try:
            system_group_by_unit_id[UUID(raw_unit_id)] = conversation
        except ValueError:
            continue
    system_group_ids = [conversation.id for conversation in system_group_by_unit_id.values()]
    linked_system_group_ids: set[UUID] = set()
    if system_group_ids:
        linked_system_group_ids.update(
            (
                await db.scalars(
                    select(ConversationMember.conversation_id).where(
                        ConversationMember.conversation_id.in_(system_group_ids)
                    )
                )
            ).all()
        )
        linked_system_group_ids.update(
            (
                await db.scalars(
                    select(ConversationInvite.conversation_id).where(
                        ConversationInvite.conversation_id.in_(system_group_ids)
                    )
                )
            ).all()
        )
        linked_system_group_ids.update(
            (
                await db.scalars(
                    select(Message.conversation_id).where(
                        Message.conversation_id.in_(system_group_ids)
                    )
                )
            ).all()
        )

    referenced: dict[UUID, set[str]] = {unit_id: set() for unit_id in leftover_ids}
    for user in users:
        for unit_id in (user.college_id, user.school_id, user.department_id):
            if unit_id in referenced:
                referenced[unit_id].add("member affiliation")
    for unit_id in _referenced_unit_ids(announcements):
        if unit_id in referenced:
            referenced[unit_id].add("announcement audience")
    for unit_id in _referenced_unit_ids(documents):
        if unit_id in referenced:
            referenced[unit_id].add("document audience")
    for unit_id, conversation in system_group_by_unit_id.items():
        if unit_id in referenced and conversation.id in linked_system_group_ids:
            referenced[unit_id].add("system chat history")

    protected = {unit_id for unit_id, reasons in referenced.items() if reasons}
    changed = True
    while changed:
        changed = False
        for item in units:
            if item.parent_id not in leftover_ids:
                continue
            child_is_kept = item.id not in leftover_ids or item.id in protected
            if child_is_kept and item.parent_id not in protected:
                protected.add(item.parent_id)
                referenced[item.parent_id].add("child organization unit")
                changed = True

    prunable = leftover_ids - protected
    for item in sorted(
        (units_by_id[unit_id] for unit_id in prunable),
        key=lambda value: _unit_depth(value, units_by_id),
        reverse=True,
    ):
        empty_system_group = system_group_by_unit_id.get(item.id)
        if empty_system_group is not None:
            await db.delete(empty_system_group)
        await db.delete(item)
    await db.flush()

    retained = tuple(
        f"{units_by_id[unit_id].unit_type} '{units_by_id[unit_id].name}' — "
        + ", ".join(sorted(referenced[unit_id]))
        for unit_id in sorted(
            protected,
            key=lambda value: units_by_id[value].name.casefold(),
        )
        if unit_id in units_by_id
    )
    return len(prunable), retained


async def sync_member_roster(
    db: AsyncSession,
    rows: Sequence[RosterRow],
    *,
    created_by_id: UUID | None = None,
) -> MemberRosterSyncResult:
    structure = organization_structure_from_rows(rows)
    if not structure:
        raise ApiError(
            422,
            "roster_organization_missing",
            "The roster does not contain college and school names",
        )
    seed_result = await seed_organization(db, structure)
    units = list((await db.scalars(select(OrganizationUnit))).all())
    users = list((await db.scalars(select(User))).all())
    users_by_email = {user.email: user for user in users}
    users_by_staff = {user.staff_id.casefold(): user for user in users if user.staff_id}
    member_role = await db.scalar(select(Role).where(Role.key == "member"))
    if member_role is None:
        raise RuntimeError("The member role is missing; run seed before roster sync")

    conflicts: list[str] = []
    row_issues: list[str] = []
    seen_emails: set[str] = set()
    seen_staff_ids: set[str] = set()
    matched_user_ids: set[UUID] = set()
    affected_users: dict[UUID, User] = {}
    members_relinked = 0
    titles_filled = 0
    ranks_filled = 0
    staff_ids_filled = 0
    members_created = 0
    members_unchanged = 0
    now = datetime.now(UTC)

    for row in rows:
        label = _row_label(row)
        email = normalize_email(row.email) if row.email else ""
        staff_key = row.staff_id.casefold() if row.staff_id else ""
        if email and email in seen_emails:
            row_issues.append(f"{label}: duplicate email in the roster")
            continue
        if staff_key and staff_key in seen_staff_ids:
            row_issues.append(f"{label}: duplicate staff ID in the roster")
            continue
        if email:
            seen_emails.add(email)
        if staff_key:
            seen_staff_ids.add(staff_key)

        by_email = users_by_email.get(email) if email else None
        by_staff = users_by_staff.get(staff_key) if staff_key else None
        user: User | None = None
        if by_email is not None and by_staff is not None and by_email.id != by_staff.id:
            conflicts.append(
                f"{label}: email belongs to {by_email.email} but staff ID belongs to "
                f"{by_staff.email}"
            )
            continue
        if by_email is not None:
            if by_email.staff_id and staff_key and by_email.staff_id.casefold() != staff_key:
                conflicts.append(f"{label}: email matches {by_email.email} but staff IDs differ")
                continue
            user = by_email
        elif by_staff is not None:
            if email and by_staff.email != email:
                conflicts.append(
                    f"{label}: staff ID matches {by_staff.email} but the roster email differs"
                )
                continue
            user = by_staff
        elif not email or not _valid_email(row.email):
            row_issues.append(f"{label}: a valid email is required")
            continue
        elif not row.staff_id:
            row_issues.append(f"{label}: staff ID is required to create a member")
            continue
        elif not row.other_name or not row.surname:
            row_issues.append(f"{label}: other name and surname are required to create a member")
            continue

        college = (
            _find_unit(units, unit_type="college", name=row.college, parent_id=None)
            if row.college
            else None
        )
        school = (
            _find_unit(
                units,
                unit_type="school",
                name=row.school,
                parent_id=college.id if college else None,
            )
            if row.school and college is not None
            else None
        )
        department = (
            _find_unit(
                units,
                unit_type="department",
                name=row.department,
                parent_id=school.id if school else None,
            )
            if row.department and school is not None
            else None
        )
        if not row.college or not row.school:
            row_issues.append(f"{label}: college and school are required to set affiliation")
        elif row.college and college is None:
            row_issues.append(f"{label}: college '{row.college}' was not found")
        elif row.school and school is None:
            row_issues.append(f"{label}: school '{row.school}' was not found")
        elif row.department and department is None:
            row_issues.append(f"{label}: department '{row.department}' was not found")

        next_college_id = college.id if college else None
        next_school_id = school.id if school else None
        next_department_id = department.id if department else None
        can_assign_org = college is not None and school is not None

        if user is None:
            created = User(
                id=new_id(),
                email=email,
                staff_id=row.staff_id,
                password_hash=hash_password(row.staff_id),
                must_change_password=True,
                email_verified=False,
                status="active",
                title=row.title,
                other_name=row.other_name,
                surname=row.surname,
                gender=row.gender or None,
                academic_rank=row.academic_rank or None,
                phone_number=row.phone_number or None,
                college_id=next_college_id if can_assign_org else None,
                school_id=next_school_id if can_assign_org else None,
                department_id=next_department_id if can_assign_org else None,
            )
            db.add(created)
            await db.flush()
            db.add(
                UserRole(
                    id=new_id(),
                    user_id=created.id,
                    role_id=member_role.id,
                    assigned_by_id=created_by_id,
                    assigned_at=now,
                )
            )
            users_by_email[created.email] = created
            users_by_staff[row.staff_id.casefold()] = created
            users.append(created)
            matched_user_ids.add(created.id)
            affected_users[created.id] = created
            members_created += 1
            continue

        matched_user_ids.add(user.id)
        changed = False
        if can_assign_org and (
            user.college_id,
            user.school_id,
            user.department_id,
        ) != (next_college_id, next_school_id, next_department_id):
            user.college_id = next_college_id
            user.school_id = next_school_id
            user.department_id = next_department_id
            members_relinked += 1
            changed = True
        if _blank(user.title) and row.title:
            user.title = row.title
            titles_filled += 1
            changed = True
        if _blank(user.academic_rank) and row.academic_rank:
            user.academic_rank = row.academic_rank
            ranks_filled += 1
            changed = True
        if _blank(user.staff_id) and row.staff_id:
            user.staff_id = row.staff_id
            users_by_staff[row.staff_id.casefold()] = user
            staff_ids_filled += 1
            changed = True
        if changed:
            affected_users[user.id] = user
        else:
            members_unchanged += 1

    await db.flush()
    users = list((await db.scalars(select(User))).all())
    pruned, retained = await _prune_leftover_units(
        db,
        canonical_ids=seed_result.canonical_ids,
        users=users,
    )
    chat = await sync_system_chat_groups(
        db,
        list(affected_users.values()),
        created_by_id=created_by_id,
    )
    unmatched = sum(1 for user in users if user.id not in matched_user_ids)
    return MemberRosterSyncResult(
        rows_total=len(rows),
        units_created=seed_result.created,
        units_updated=seed_result.updated,
        members_relinked=members_relinked,
        titles_filled=titles_filled,
        ranks_filled=ranks_filled,
        staff_ids_filled=staff_ids_filled,
        members_created=members_created,
        members_unchanged=members_unchanged,
        unmatched_existing_members=unmatched,
        units_pruned=pruned,
        units_retained=len(retained),
        chat_groups_created=chat.groups_created,
        chat_memberships_added=chat.memberships_added,
        chat_memberships_removed=chat.memberships_removed,
        conflicts=tuple(conflicts),
        row_issues=tuple(row_issues),
        retained_units=retained,
        affected_user_ids=tuple(affected_users),
    )
