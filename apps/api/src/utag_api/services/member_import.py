import csv
import io
import re
import zipfile
from collections.abc import Iterable
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from utag_api.errors import ApiError

MAX_IMPORT_BYTES = 8_000_000
MAX_IMPORT_ROWS = 2_000
MAX_XLSX_EXPANDED_BYTES = 50_000_000

HEADER_ALIASES = {
    "academic_rank": "academic_rank",
    "college": "college",
    "college_name": "college",
    "department": "department",
    "department_name": "department",
    "email": "email",
    "email_address": "email",
    "e_mail": "email",
    "faculty": "school",
    "faculty_school": "school",
    "first_name": "other_name",
    "firstname": "other_name",
    "gender": "gender",
    "groups": "roles",
    "last_name": "surname",
    "lastname": "surname",
    "other_name": "other_name",
    "other_names": "other_name",
    "phone": "phone_number",
    "phone_number": "phone_number",
    "rank": "academic_rank",
    "role": "roles",
    "roles": "roles",
    "school": "school",
    "school_name": "school",
    "staff_id": "staff_id",
    "staffid": "staff_id",
    "surname": "surname",
    "telephone": "phone_number",
    "title": "title",
}


def normalized_header(value: object) -> str:
    text = re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().casefold()).strip("_")
    return HEADER_ALIASES.get(text, text)


def clean_cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalized_row(headers: list[str], values: Iterable[object]) -> dict[str, str | list[str]]:
    row: dict[str, str | list[str]] = {
        headers[index]: clean_cell(value)
        for index, value in enumerate(values)
        if index < len(headers)
    }
    roles = re.split(r"[,;|]", str(row.get("roles", "member")))
    row["roles"] = sorted({item.strip().casefold() for item in roles if item.strip()} or {"member"})
    return row


def csv_rows(data: bytes) -> list[dict[str, str | list[str]]]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ApiError(422, "import_encoding_invalid", "CSV files must use UTF-8 encoding") from exc
    reader = csv.reader(io.StringIO(text))
    try:
        raw_headers = next(reader)
    except StopIteration as exc:
        raise ApiError(422, "import_empty", "The import file is empty") from exc
    headers = [normalized_header(value) for value in raw_headers]
    return [
        normalized_row(headers, values)
        for values in reader
        if any(clean_cell(value) for value in values)
    ]


def verify_xlsx_archive(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            expanded = sum(item.file_size for item in archive.infolist())
            if expanded > MAX_XLSX_EXPANDED_BYTES:
                raise ApiError(413, "import_too_large", "The expanded workbook is too large")
    except zipfile.BadZipFile as exc:
        raise ApiError(422, "workbook_invalid", "The Excel workbook is not valid") from exc


def xlsx_rows(data: bytes) -> list[dict[str, str | list[str]]]:
    verify_xlsx_archive(data)
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (OSError, ValueError, KeyError, zipfile.BadZipFile) as exc:
        raise ApiError(422, "workbook_invalid", "The Excel workbook could not be read") from exc
    try:
        worksheet = workbook.active
        values = worksheet.iter_rows(values_only=True)
        raw_headers = next(values, None)
        if raw_headers is None:
            raise ApiError(422, "import_empty", "The import file is empty")
        headers = [normalized_header(value) for value in raw_headers]
        return [
            normalized_row(headers, row)
            for row in values
            if any(clean_cell(value) for value in row)
        ]
    finally:
        workbook.close()


def parse_member_import(filename: str, data: bytes) -> list[dict[str, Any]]:
    if not data:
        raise ApiError(422, "import_empty", "The import file is empty")
    if len(data) > MAX_IMPORT_BYTES:
        raise ApiError(413, "import_too_large", "Member import files must be 8 MB or smaller")
    suffix = Path(filename).suffix.casefold()
    if suffix == ".csv":
        rows = csv_rows(data)
    elif suffix == ".xlsx":
        rows = xlsx_rows(data)
    else:
        raise ApiError(415, "import_type_invalid", "Choose a CSV or XLSX member file")
    if len(rows) > MAX_IMPORT_ROWS:
        raise ApiError(413, "import_row_limit", f"Imports are limited to {MAX_IMPORT_ROWS} rows")
    return rows
