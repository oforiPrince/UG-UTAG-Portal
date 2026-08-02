import csv
import io
from collections.abc import Sequence
from datetime import UTC, datetime, timedelta
from typing import Annotated
from uuid import UUID

from fastapi import APIRouter, Depends, File, Form, Query, Request, UploadFile
from fastapi.responses import Response, StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import ValidationError
from sqlalchemy import delete, func, or_, select, update

from utag_api.database import new_id
from utag_api.dependencies import (
    DbSession,
    Principal,
    event_context,
    require_mutation_permissions,
    require_permissions,
)
from utag_api.errors import ApiError
from utag_api.models import (
    AccountToken,
    BackgroundJob,
    OrganizationUnit,
    Permission,
    Role,
    Session,
    User,
    UserPermissionGrant,
    UserRole,
)
from utag_api.permissions import NON_DELEGABLE_DIRECT_PERMISSIONS
from utag_api.schemas.common import MessageResponse, Page
from utag_api.schemas.domain import (
    MemberCreate,
    MemberImportIssue,
    MemberImportResult,
    MemberLifecycleRequest,
    MemberPermissionUpdate,
    MemberUpdate,
    MemberView,
    PermissionOption,
)
from utag_api.security import encrypt_text, hash_password, new_token, normalize_email, token_digest
from utag_api.services.events import enqueue_task, record_change
from utag_api.services.identity import (
    UserAccess,
    load_user_access,
    load_user_access_map,
    require_public_profile_image,
)
from utag_api.services.member_import import MAX_IMPORT_BYTES, parse_member_import
from utag_api.services.organization import (
    organization_assignment_errors,
    validate_organization_assignment,
)
from utag_api.services.query import paginate
from utag_api.services.system_chat_groups import sync_system_chat_groups

router = APIRouter(prefix="/members", tags=["members"])


def normalized_organization_name(value: object) -> str:
    return " ".join(str(value or "").split()).casefold()


def resolve_import_organization(
    raw: dict[str, object],
    units: dict[UUID, OrganizationUnit],
) -> tuple[dict[str, object], list[tuple[str, str]]]:
    values = dict(raw)
    issues: list[tuple[str, str]] = []

    def resolve_unit(
        field: str,
        unit_type: str,
        parent_id: UUID | None,
    ) -> OrganizationUnit | None:
        raw_id = values.get(f"{field}_id")
        if raw_id:
            try:
                unit_id = UUID(str(raw_id))
            except ValueError:
                issues.append((f"{field}_id", f"Select a valid {unit_type}"))
                return None
            unit = units.get(unit_id)
            if (
                unit is None
                or unit.unit_type != unit_type
                or not unit.is_active
                or unit.parent_id != parent_id
            ):
                issues.append((f"{field}_id", f"Select a valid active {unit_type}"))
                return None
            return unit

        name = normalized_organization_name(values.get(field))
        if not name:
            issues.append((field, f"Enter the member's {unit_type}"))
            return None
        matches = [
            unit
            for unit in units.values()
            if unit.unit_type == unit_type
            and unit.is_active
            and unit.parent_id == parent_id
            and normalized_organization_name(unit.name) == name
        ]
        if len(matches) != 1:
            issues.append(
                (
                    field,
                    f"{unit_type.title()} was not found in the organization reference",
                )
            )
            return None
        return matches[0]

    college = resolve_unit("college", "college", None)
    school = resolve_unit("school", "school", college.id if college else None)
    department = resolve_unit(
        "department",
        "department",
        school.id if school else None,
    )
    if college:
        values["college_id"] = college.id
    if school:
        values["school_id"] = school.id
    if department:
        values["department_id"] = department.id
    for field in ("college", "school", "department"):
        values.pop(field, None)
    return values, issues


def import_chat_group_titles(
    candidate: MemberCreate,
    units: dict[UUID, OrganizationUnit],
) -> list[str]:
    titles = ["UTAG UG"]
    if candidate.school_id and candidate.school_id in units:
        titles.append(f"School: {units[candidate.school_id].name}")
    if candidate.department_id and candidate.department_id in units:
        titles.append(f"Department: {units[candidate.department_id].name}")
    return titles


async def revoke_member_sessions(db: DbSession, user_id: UUID, now: datetime) -> None:
    await db.execute(
        update(Session)
        .where(Session.user_id == user_id, Session.revoked_at.is_(None))
        .values(revoked_at=now)
    )


async def invalidate_member_access_tokens(db: DbSession, user_id: UUID, now: datetime) -> None:
    await db.execute(
        update(AccountToken)
        .where(
            AccountToken.user_id == user_id,
            AccountToken.used_at.is_(None),
        )
        .values(used_at=now)
    )


async def ensure_privileged_target_permission(
    db: DbSession, principal: Principal, user_id: UUID
) -> None:
    if "members.roles" in principal.permissions:
        return
    administrator_role = await db.scalar(
        select(Role.key)
        .join(UserRole, UserRole.role_id == Role.id)
        .where(UserRole.user_id == user_id, Role.key == "administrator")
    )
    if administrator_role:
        raise ApiError(
            403,
            "administrator_management_denied",
            "Only an administrator can manage another administrator account",
        )


async def active_administrator_ids(db: DbSession) -> list[UUID]:
    return list(
        (
            await db.scalars(
                select(User.id)
                .join(UserRole, UserRole.user_id == User.id)
                .join(Role, Role.id == UserRole.role_id)
                .where(User.status == "active", Role.key == "administrator")
                .with_for_update()
            )
        ).all()
    )


async def protect_administrator_removal(
    db: DbSession,
    *,
    target: User,
    principal: Principal,
    replacement_roles: set[str] | None = None,
) -> None:
    target_is_administrator = bool(
        await db.scalar(
            select(UserRole.user_id)
            .join(Role, Role.id == UserRole.role_id)
            .where(UserRole.user_id == target.id, Role.key == "administrator")
        )
    )
    if not target_is_administrator:
        return
    if replacement_roles is not None and "administrator" in replacement_roles:
        return
    if target.id == principal.user.id:
        raise ApiError(
            409,
            "self_administrator_change_denied",
            "Another administrator must change your administrator access",
        )
    active_ids = await active_administrator_ids(db)
    if target.status == "active" and active_ids == [target.id]:
        raise ApiError(
            409,
            "last_administrator_required",
            "At least one active administrator must remain",
        )


async def organization_name_map(db: DbSession, users: Sequence[User]) -> dict[UUID, str]:
    unit_ids = {
        unit_id
        for user in users
        for unit_id in (user.school_id, user.college_id, user.department_id)
        if unit_id is not None
    }
    if not unit_ids:
        return {}
    units = (
        await db.scalars(select(OrganizationUnit).where(OrganizationUnit.id.in_(unit_ids)))
    ).all()
    return {unit.id: unit.name for unit in units}


def member_view(
    user: User,
    access: UserAccess,
    organization_names: dict[UUID, str] | None = None,
) -> MemberView:
    names = organization_names or {}
    return MemberView(
        id=user.id,
        email=user.email,
        staff_id=user.staff_id,
        title=user.title,
        other_name=user.other_name,
        surname=user.surname,
        gender=user.gender,
        academic_rank=user.academic_rank,
        phone_number=user.phone_number,
        profile_media_id=getattr(user, "profile_media_id", None),
        school_id=user.school_id,
        college_id=user.college_id,
        department_id=user.department_id,
        full_name=user.full_name,
        status=user.status,
        email_verified=user.email_verified,
        must_change_password=user.must_change_password,
        roles=sorted(access.roles),
        role_permissions=sorted(access.role_permissions),
        extra_permissions=sorted(access.extra_permissions),
        effective_permissions=sorted(access.effective_permissions),
        created_at=user.created_at,
        last_login_at=getattr(user, "last_login_at", None),
        school_name=names.get(user.school_id) if user.school_id else None,
        college_name=names.get(user.college_id) if user.college_id else None,
        department_name=names.get(user.department_id) if user.department_id else None,
    )


def csv_safe(value: object) -> object:
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@", "\t", "\r")):
        return f"'{value}"
    return value


@router.get("/permission-options", response_model=list[PermissionOption])
async def list_permission_options(
    db: DbSession,
    principal: Annotated[Principal, Depends(require_permissions("members.permissions"))],
) -> list[PermissionOption]:
    permissions = (
        await db.scalars(
            select(Permission)
            .where(Permission.key.not_in(NON_DELEGABLE_DIRECT_PERMISSIONS))
            .order_by(Permission.key)
        )
    ).all()
    return [
        PermissionOption(key=permission.key, description=permission.description)
        for permission in permissions
    ]


@router.get("", response_model=Page[MemberView])
async def list_members(
    db: DbSession,
    principal: Annotated[Principal, Depends(require_permissions("members.view"))],
    page: Annotated[int, Query(ge=1)] = 1,
    page_size: Annotated[int, Query(ge=1, le=100)] = 25,
    q: str | None = None,
    status: str | None = None,
    unit_id: UUID | None = None,
    role: str | None = None,
) -> Page[MemberView]:
    statement = select(User).order_by(User.surname, User.other_name)
    if q:
        pattern = f"%{q.strip()}%"
        statement = statement.where(
            or_(
                User.email.ilike(pattern),
                User.staff_id.ilike(pattern),
                User.other_name.ilike(pattern),
                User.surname.ilike(pattern),
            )
        )
    if status:
        statement = statement.where(User.status == status)
    if unit_id:
        statement = statement.where(
            or_(
                User.school_id == unit_id,
                User.college_id == unit_id,
                User.department_id == unit_id,
            )
        )
    if role:
        statement = (
            statement.join(UserRole, UserRole.user_id == User.id)
            .join(Role, Role.id == UserRole.role_id)
            .where(Role.key == role)
        )
    result = await paginate(db, statement, page=page, page_size=page_size)
    access = await load_user_access_map(db, [item.id for item in result.items])
    organization_names = await organization_name_map(db, result.items)
    return Page[MemberView](
        items=[member_view(item, access[item.id], organization_names) for item in result.items],
        page=result.page,
        page_size=result.page_size,
        total=result.total,
        pages=result.pages,
    )


@router.get("/{user_id:uuid}", response_model=MemberView)
async def get_member(
    user_id: UUID,
    db: DbSession,
    principal: Annotated[Principal, Depends(require_permissions("members.view"))],
) -> MemberView:
    user = await db.get(User, user_id)
    if user is None:
        raise ApiError(404, "member_not_found", "Member not found")
    return member_view(
        user,
        await load_user_access(db, user.id),
        await organization_name_map(db, [user]),
    )


@router.post("", response_model=MemberView, status_code=201)
async def create_member(
    payload: MemberCreate,
    request: Request,
    db: DbSession,
    principal: Annotated[Principal, Depends(require_mutation_permissions("members.create"))],
) -> MemberView:
    email = normalize_email(str(payload.email))
    if await db.scalar(select(User.id).where(User.email == email)):
        raise ApiError(409, "email_exists", "A member already uses this email")
    if payload.profile_media_id is not None:
        await require_public_profile_image(db, payload.profile_media_id)
    await validate_organization_assignment(
        db,
        school_id=payload.school_id,
        college_id=payload.college_id,
        department_id=payload.department_id,
    )
    requested_roles = (await db.scalars(select(Role).where(Role.key.in_(set(payload.roles))))).all()
    if len(requested_roles) != len(set(payload.roles)):
        raise ApiError(422, "invalid_role", "One or more roles do not exist")
    if set(payload.roles) != {"member"} and "members.roles" not in principal.permissions:
        raise ApiError(
            403,
            "role_assignment_denied",
            "You do not have permission to assign privileged roles",
        )
    user = User(
        id=new_id(),
        email=email,
        staff_id=payload.staff_id,
        password_hash=hash_password(new_token()),
        must_change_password=True,
        email_verified=False,
        status="invited" if payload.send_invitation else "active",
        title=payload.title,
        other_name=payload.other_name,
        surname=payload.surname,
        gender=payload.gender,
        academic_rank=payload.academic_rank,
        phone_number=payload.phone_number,
        profile_media_id=payload.profile_media_id,
        school_id=payload.school_id,
        college_id=payload.college_id,
        department_id=payload.department_id,
    )
    db.add(user)
    await db.flush()
    now = datetime.now(UTC)
    for role in requested_roles:
        db.add(
            UserRole(
                user_id=user.id,
                role_id=role.id,
                assigned_by_id=principal.user.id,
                assigned_at=now,
            )
        )
    invitation_token: str | None = None
    if payload.send_invitation:
        invitation_token = new_token()
        db.add(
            AccountToken(
                user_id=user.id,
                kind="invitation",
                token_hash=token_digest(invitation_token),
                created_at=now,
                expires_at=now + timedelta(days=7),
            )
        )
        encrypted_token = encrypt_text(invitation_token)
        if encrypted_token is None:
            raise RuntimeError("Invitation token encryption failed")
        enqueue_task(
            db,
            task_name="utag.email.invitation",
            aggregate_type="user",
            aggregate_id=user.id,
            args=[str(user.id), encrypted_token.decode()],
            queue="communications",
        )
    await sync_system_chat_groups(
        db,
        [user],
        created_by_id=principal.user.id,
    )
    record_change(
        db,
        context=event_context(request, principal),
        action="member.created",
        resource_type="user",
        resource_id=user.id,
        topic="members",
        payload={"user_id": str(user.id), "status": user.status},
    )
    await db.commit()
    return member_view(
        user,
        await load_user_access(db, user.id),
        await organization_name_map(db, [user]),
    )


@router.put("/{user_id:uuid}/permissions", response_model=MemberView)
async def update_member_permissions(
    user_id: UUID,
    payload: MemberPermissionUpdate,
    request: Request,
    db: DbSession,
    principal: Annotated[Principal, Depends(require_mutation_permissions("members.permissions"))],
) -> MemberView:
    user = await db.get(User, user_id)
    if user is None:
        raise ApiError(404, "member_not_found", "Member not found")
    await ensure_privileged_target_permission(db, principal, user.id)

    requested_keys = set(payload.permissions)
    non_delegable = sorted(requested_keys & NON_DELEGABLE_DIRECT_PERMISSIONS)
    if non_delegable:
        raise ApiError(
            422,
            "permission_not_delegable",
            "Role and permission administration access can only be assigned through a role",
            details={"permissions": non_delegable},
        )
    permissions = (
        await db.scalars(select(Permission).where(Permission.key.in_(requested_keys)))
    ).all()
    permission_by_key = {permission.key: permission for permission in permissions}
    unknown = sorted(requested_keys - permission_by_key.keys())
    if unknown:
        raise ApiError(
            422,
            "invalid_permission",
            "One or more permissions do not exist",
            details={"permissions": unknown},
        )

    previous_access = await load_user_access(db, user.id)
    previous_permissions = sorted(previous_access.extra_permissions)
    await db.execute(delete(UserPermissionGrant).where(UserPermissionGrant.user_id == user.id))
    now = datetime.now(UTC)
    for permission_key in sorted(requested_keys):
        db.add(
            UserPermissionGrant(
                id=new_id(),
                user_id=user.id,
                permission_id=permission_by_key[permission_key].id,
                granted_by_id=principal.user.id,
                granted_at=now,
            )
        )
    record_change(
        db,
        context=event_context(request, principal),
        action="member.permissions.updated",
        resource_type="user",
        resource_id=user.id,
        topic="members",
        payload={"user_id": str(user.id), "permissions": sorted(requested_keys)},
        changes={
            "extra_permissions": {
                "from": previous_permissions,
                "to": sorted(requested_keys),
            }
        },
    )
    await db.commit()
    return member_view(
        user,
        await load_user_access(db, user.id),
        await organization_name_map(db, [user]),
    )


@router.post("/import", response_model=MemberImportResult)
async def import_members(
    request: Request,
    db: DbSession,
    principal: Annotated[Principal, Depends(require_mutation_permissions("members.create"))],
    file: Annotated[UploadFile, File()],
    dry_run: Annotated[bool, Form()] = True,
    force_role: Annotated[str | None, Form()] = None,
) -> MemberImportResult:
    data = await file.read(MAX_IMPORT_BYTES + 1)
    raw_rows = parse_member_import(file.filename or "members.csv", data)
    issues: list[MemberImportIssue] = []
    candidates: list[tuple[int, MemberCreate]] = []
    seen_emails: set[str] = set()
    seen_staff_ids: set[str] = set()
    duplicate_rows = 0
    units = {
        unit.id: unit
        for unit in (
            await db.scalars(select(OrganizationUnit).where(OrganizationUnit.is_active.is_(True)))
        ).all()
    }

    for row_number, raw in enumerate(raw_rows, start=2):
        resolved, organization_issues = resolve_import_organization(raw, units)
        if organization_issues:
            issues.extend(
                MemberImportIssue(row=row_number, field=field, message=message)
                for field, message in organization_issues
            )
            continue
        try:
            roles = [force_role] if force_role else resolved.get("roles") or ["member"]
            candidate = MemberCreate.model_validate(
                {**resolved, "send_invitation": False, "roles": roles}
            )
        except ValidationError as exc:
            for error in exc.errors(include_url=False):
                location = ".".join(str(item) for item in error["loc"])
                issues.append(
                    MemberImportIssue(
                        row=row_number,
                        field=location or None,
                        message=str(error["msg"]),
                    )
                )
            continue
        if not candidate.staff_id:
            issues.append(
                MemberImportIssue(
                    row=row_number,
                    field="staff_id",
                    message="Staff ID is required and becomes the temporary password",
                )
            )
            continue
        normalized = normalize_email(str(candidate.email))
        if normalized in seen_emails:
            duplicate_rows += 1
            issues.append(
                MemberImportIssue(
                    row=row_number,
                    field="email",
                    message="This email appears more than once in the file",
                )
            )
            continue
        seen_emails.add(normalized)
        normalized_staff_id = candidate.staff_id.casefold()
        if normalized_staff_id in seen_staff_ids:
            duplicate_rows += 1
            issues.append(
                MemberImportIssue(
                    row=row_number,
                    field="staff_id",
                    message="This staff ID appears more than once in the file",
                )
            )
            continue
        seen_staff_ids.add(normalized_staff_id)
        candidates.append((row_number, candidate))

    candidate_emails = [normalize_email(str(candidate.email)) for _, candidate in candidates]
    existing_emails = set(
        (await db.scalars(select(User.email).where(User.email.in_(candidate_emails)))).all()
    )
    existing_staff_ids = {
        str(staff_id).casefold(): email
        for staff_id, email in (
            await db.execute(
                select(User.staff_id, User.email).where(
                    User.staff_id.is_not(None),
                    func.lower(User.staff_id).in_(seen_staff_ids),
                )
            )
        ).all()
        if staff_id
    }
    valid_roles = set((await db.scalars(select(Role.key))).all())
    requested_role_keys = {role for _, candidate in candidates for role in candidate.roles}
    if (
        requested_role_keys
        and requested_role_keys != {"member"}
        and "members.roles" not in principal.permissions
    ):
        raise ApiError(
            403,
            "role_assignment_denied",
            "You do not have permission to import privileged roles",
        )
    accepted: list[tuple[int, MemberCreate]] = []
    for row_number, candidate in candidates:
        email = normalize_email(str(candidate.email))
        row_has_error = False
        if email in existing_emails:
            duplicate_rows += 1
            row_has_error = True
            issues.append(
                MemberImportIssue(
                    row=row_number,
                    field="email",
                    message="A member already uses this email",
                )
            )
        staff_owner = existing_staff_ids.get((candidate.staff_id or "").casefold())
        if staff_owner is not None:
            duplicate_rows += 1
            row_has_error = True
            issues.append(
                MemberImportIssue(
                    row=row_number,
                    field="staff_id",
                    message=f"Another member already uses this staff ID ({staff_owner})",
                )
            )
        unknown_roles = sorted(set(candidate.roles) - valid_roles)
        if unknown_roles:
            row_has_error = True
            issues.append(
                MemberImportIssue(
                    row=row_number,
                    field="roles",
                    message=f"Unknown role(s): {', '.join(unknown_roles)}",
                )
            )
        for field, expected_type in (
            ("school_id", "school"),
            ("college_id", "college"),
            ("department_id", "department"),
        ):
            unit_id = getattr(candidate, field)
            if unit_id is None:
                continue
            unit = units.get(unit_id)
            if unit is None or unit.unit_type != expected_type or not unit.is_active:
                row_has_error = True
                issues.append(
                    MemberImportIssue(
                        row=row_number,
                        field=field,
                        message=f"Select a valid active {expected_type}",
                    )
                )
        hierarchy_errors = organization_assignment_errors(
            units,
            school_id=candidate.school_id,
            college_id=candidate.college_id,
            department_id=candidate.department_id,
        )
        for field, message in hierarchy_errors:
            if any(
                issue.row == row_number and issue.field == field and issue.message == message
                for issue in issues
            ):
                continue
            row_has_error = True
            issues.append(
                MemberImportIssue(
                    row=row_number,
                    field=field,
                    message=message,
                )
            )
        if not row_has_error:
            accepted.append((row_number, candidate))

    invalid_rows = len({issue.row for issue in issues})
    preview = [
        {
            "row": row_number,
            "email": str(candidate.email),
            "staff_id": candidate.staff_id,
            "full_name": " ".join(
                value
                for value in (candidate.title, candidate.other_name, candidate.surname)
                if value
            ),
            "academic_rank": candidate.academic_rank,
            "roles": candidate.roles,
            "chat_groups": import_chat_group_titles(candidate, units),
        }
        for row_number, candidate in accepted[:50]
    ]
    result = MemberImportResult(
        dry_run=dry_run,
        total_rows=len(raw_rows),
        valid_rows=len(accepted),
        invalid_rows=invalid_rows,
        imported_rows=0,
        duplicate_rows=duplicate_rows,
        issues=issues,
        preview=preview,
        preview_truncated=len(accepted) > len(preview),
    )
    if dry_run or issues:
        return result

    roles_by_key = {role.key: role for role in (await db.scalars(select(Role))).all()}
    now = datetime.now(UTC)
    imported_users: list[User] = []
    for _, candidate in accepted:
        if candidate.staff_id is None:
            raise RuntimeError("Validated member import is missing a staff ID")
        user = User(
            id=new_id(),
            email=normalize_email(str(candidate.email)),
            staff_id=candidate.staff_id,
            password_hash=hash_password(candidate.staff_id),
            must_change_password=True,
            email_verified=False,
            status="active",
            title=candidate.title,
            other_name=candidate.other_name,
            surname=candidate.surname,
            gender=candidate.gender,
            academic_rank=candidate.academic_rank,
            phone_number=candidate.phone_number,
            school_id=candidate.school_id,
            college_id=candidate.college_id,
            department_id=candidate.department_id,
        )
        db.add(user)
        imported_users.append(user)
        await db.flush()
        for role_key in candidate.roles:
            db.add(
                UserRole(
                    id=new_id(),
                    user_id=user.id,
                    role_id=roles_by_key[role_key].id,
                    assigned_by_id=principal.user.id,
                    assigned_at=now,
                )
            )
        record_change(
            db,
            context=event_context(request, principal),
            action="member.imported",
            resource_type="user",
            resource_id=user.id,
            topic="members",
            payload={"user_id": str(user.id), "status": user.status},
        )
    chat_sync = await sync_system_chat_groups(
        db,
        imported_users,
        created_by_id=principal.user.id,
    )
    result.imported_rows = len(accepted)
    result.chat_groups_created = chat_sync.groups_created
    result.chat_memberships_added = chat_sync.memberships_added
    db.add(
        BackgroundJob(
            id=new_id(),
            owner_id=principal.user.id,
            kind="members.import",
            status="completed",
            progress=100,
            input_json={"filename": file.filename, "rows": len(raw_rows)},
            result_json={
                "imported": result.imported_rows,
                "duplicates": result.duplicate_rows,
                "invalid": result.invalid_rows,
                "chat_groups_created": result.chat_groups_created,
                "chat_memberships_added": result.chat_memberships_added,
            },
        )
    )
    await db.commit()
    return result


@router.patch("/{user_id:uuid}", response_model=MemberView)
async def update_member(
    user_id: UUID,
    payload: MemberUpdate,
    request: Request,
    db: DbSession,
    principal: Annotated[Principal, Depends(require_mutation_permissions("members.update"))],
) -> MemberView:
    user = await db.get(User, user_id)
    if user is None:
        raise ApiError(404, "member_not_found", "Member not found")
    await ensure_privileged_target_permission(db, principal, user.id)
    if payload.status is not None and payload.status != user.status:
        raise ApiError(
            409,
            "member_lifecycle_action_required",
            "Use the deactivate, reactivate, or archive action to change account status",
        )
    changes = payload.model_dump(exclude_unset=True, exclude={"roles", "status"})
    await validate_organization_assignment(
        db,
        school_id=changes.get("school_id", user.school_id),
        college_id=changes.get("college_id", user.college_id),
        department_id=changes.get("department_id", user.department_id),
    )
    profile_media_id = changes.get("profile_media_id")
    if profile_media_id is not None:
        await require_public_profile_image(db, profile_media_id)
    for key, value in changes.items():
        setattr(user, key, value)
    if payload.roles is not None:
        if "members.roles" not in principal.permissions:
            raise ApiError(
                403,
                "role_assignment_denied",
                "You do not have permission to assign roles",
            )
        roles = (await db.scalars(select(Role).where(Role.key.in_(set(payload.roles))))).all()
        if len(roles) != len(set(payload.roles)):
            raise ApiError(422, "invalid_role", "One or more roles do not exist")
        await protect_administrator_removal(
            db,
            target=user,
            principal=principal,
            replacement_roles=set(payload.roles),
        )
        await db.execute(delete(UserRole).where(UserRole.user_id == user.id))
        now = datetime.now(UTC)
        for role in roles:
            db.add(
                UserRole(
                    user_id=user.id,
                    role_id=role.id,
                    assigned_by_id=principal.user.id,
                    assigned_at=now,
                )
            )
    await sync_system_chat_groups(
        db,
        [user],
        created_by_id=principal.user.id,
    )
    record_change(
        db,
        context=event_context(request, principal),
        action="member.updated",
        resource_type="user",
        resource_id=user.id,
        topic="members",
        payload={"user_id": str(user.id), "status": user.status},
        changes={key: {"to": str(value)} for key, value in changes.items()},
    )
    await db.commit()
    return member_view(
        user,
        await load_user_access(db, user.id),
        await organization_name_map(db, [user]),
    )


@router.post("/{user_id:uuid}/access-link", response_model=MessageResponse)
async def send_access_link(
    user_id: UUID,
    request: Request,
    db: DbSession,
    principal: Annotated[Principal, Depends(require_mutation_permissions("members.credentials"))],
) -> MessageResponse:
    user = await db.get(User, user_id)
    if user is None or user.status == "archived":
        raise ApiError(404, "member_not_found", "Member not found")
    await ensure_privileged_target_permission(db, principal, user.id)
    if user.status == "suspended":
        raise ApiError(
            409,
            "member_deactivated",
            "Reactivate the member before sending an access link",
        )
    now = datetime.now(UTC)
    await invalidate_member_access_tokens(db, user.id, now)
    raw_token = new_token()
    is_invitation = user.status == "invited" or not user.email_verified
    token_kind = "invitation" if is_invitation else "password_reset"
    db.add(
        AccountToken(
            user_id=user.id,
            kind=token_kind,
            token_hash=token_digest(raw_token),
            created_at=now,
            expires_at=now + (timedelta(days=7) if is_invitation else timedelta(minutes=30)),
        )
    )
    record_change(
        db,
        context=event_context(request, principal),
        action="member.access_link.sent",
        resource_type="user",
        resource_id=user.id,
        topic=f"user:{user.id}",
        payload={"user_id": str(user.id), "kind": token_kind},
    )
    encrypted_token = encrypt_text(raw_token)
    if encrypted_token is None:
        raise RuntimeError("Access token encryption failed")
    enqueue_task(
        db,
        task_name=("utag.email.invitation" if is_invitation else "utag.email.password_reset"),
        aggregate_type="user",
        aggregate_id=user.id,
        args=[str(user.id), encrypted_token.decode()],
        queue="communications",
    )
    await db.commit()
    return MessageResponse(
        message="Invitation sent" if is_invitation else "Password reset instructions sent"
    )


@router.post("/{user_id:uuid}/password-reset", response_model=MessageResponse)
async def force_password_reset(
    user_id: UUID,
    request: Request,
    db: DbSession,
    principal: Annotated[Principal, Depends(require_mutation_permissions("members.credentials"))],
) -> MessageResponse:
    if user_id == principal.user.id:
        raise ApiError(
            409,
            "self_password_reset_denied",
            "Use your profile to change your own password",
        )
    user = await db.get(User, user_id)
    if user is None or user.status != "active" or not user.email_verified:
        raise ApiError(
            409,
            "member_not_active",
            "Use an access invitation until the member has activated their account",
        )
    await ensure_privileged_target_permission(db, principal, user.id)
    now = datetime.now(UTC)
    await revoke_member_sessions(db, user.id, now)
    await invalidate_member_access_tokens(db, user.id, now)
    user.password_hash = hash_password(new_token())
    user.must_change_password = True
    raw_token = new_token()
    db.add(
        AccountToken(
            user_id=user.id,
            kind="password_reset",
            token_hash=token_digest(raw_token),
            created_at=now,
            expires_at=now + timedelta(minutes=30),
        )
    )
    record_change(
        db,
        context=event_context(request, principal),
        action="member.password_reset.required",
        resource_type="user",
        resource_id=user.id,
        topic=f"user:{user.id}",
        payload={"user_id": str(user.id)},
        reason="Administrator required a password reset",
    )
    encrypted_token = encrypt_text(raw_token)
    if encrypted_token is None:
        raise RuntimeError("Password reset token encryption failed")
    enqueue_task(
        db,
        task_name="utag.email.password_reset",
        aggregate_type="user",
        aggregate_id=user.id,
        args=[str(user.id), encrypted_token.decode()],
        queue="communications",
    )
    await db.commit()
    return MessageResponse(message="Password reset link sent and existing sessions revoked")


@router.post("/{user_id:uuid}/deactivate", response_model=MessageResponse)
async def deactivate_member(
    user_id: UUID,
    payload: MemberLifecycleRequest,
    request: Request,
    db: DbSession,
    principal: Annotated[Principal, Depends(require_mutation_permissions("members.lifecycle"))],
) -> MessageResponse:
    if user_id == principal.user.id:
        raise ApiError(
            409,
            "self_deactivation_denied",
            "You cannot deactivate your own account",
        )
    user = await db.get(User, user_id)
    if user is None:
        raise ApiError(404, "member_not_found", "Member not found")
    await ensure_privileged_target_permission(db, principal, user.id)
    await protect_administrator_removal(db, target=user, principal=principal)
    if user.status == "archived":
        raise ApiError(
            409,
            "member_archived",
            "Reactivate the archived member before changing their status",
        )
    if user.status == "suspended":
        return MessageResponse(message="Member is already deactivated")
    now = datetime.now(UTC)
    user.status = "suspended"
    await revoke_member_sessions(db, user.id, now)
    await invalidate_member_access_tokens(db, user.id, now)
    await sync_system_chat_groups(
        db,
        [user],
        created_by_id=principal.user.id,
    )
    record_change(
        db,
        context=event_context(request, principal),
        action="member.deactivated",
        resource_type="user",
        resource_id=user.id,
        topic="members",
        payload={"user_id": str(user.id), "status": user.status},
        reason=payload.reason,
    )
    await db.commit()
    return MessageResponse(message="Member deactivated and signed out")


@router.post("/{user_id:uuid}/reactivate", response_model=MessageResponse)
async def reactivate_member(
    user_id: UUID,
    request: Request,
    db: DbSession,
    principal: Annotated[Principal, Depends(require_mutation_permissions("members.lifecycle"))],
) -> MessageResponse:
    user = await db.get(User, user_id)
    if user is None:
        raise ApiError(404, "member_not_found", "Member not found")
    await ensure_privileged_target_permission(db, principal, user.id)
    if user.status == "active":
        return MessageResponse(message="Member is already active")
    if user.status not in {"suspended", "archived"}:
        raise ApiError(
            409,
            "member_not_reactivatable",
            "Invited members must accept or receive a new access invitation",
        )
    previous_status = user.status
    user.status = "active"
    await sync_system_chat_groups(
        db,
        [user],
        created_by_id=principal.user.id,
    )
    record_change(
        db,
        context=event_context(request, principal),
        action="member.reactivated",
        resource_type="user",
        resource_id=user.id,
        topic="members",
        payload={"user_id": str(user.id), "status": user.status},
        changes={"status": {"from": previous_status, "to": user.status}},
    )
    await db.commit()
    return MessageResponse(message="Member reactivated")


@router.post("/{user_id:uuid}/sessions/revoke", response_model=MessageResponse)
async def revoke_member_sessions_action(
    user_id: UUID,
    request: Request,
    db: DbSession,
    principal: Annotated[Principal, Depends(require_mutation_permissions("members.credentials"))],
) -> MessageResponse:
    if user_id == principal.user.id:
        raise ApiError(
            409,
            "self_session_revoke_denied",
            "Manage your own sessions from your profile",
        )
    user = await db.get(User, user_id)
    if user is None:
        raise ApiError(404, "member_not_found", "Member not found")
    await ensure_privileged_target_permission(db, principal, user.id)
    now = datetime.now(UTC)
    await revoke_member_sessions(db, user.id, now)
    record_change(
        db,
        context=event_context(request, principal),
        action="member.sessions.revoked",
        resource_type="user",
        resource_id=user.id,
        topic=f"user:{user.id}",
        payload={"user_id": str(user.id)},
        reason="Administrator signed the member out of all devices",
    )
    await db.commit()
    return MessageResponse(message="Member signed out of all devices")


@router.delete("/{user_id:uuid}", response_model=MessageResponse)
async def archive_member(
    user_id: UUID,
    request: Request,
    db: DbSession,
    principal: Annotated[Principal, Depends(require_mutation_permissions("members.lifecycle"))],
) -> MessageResponse:
    if user_id == principal.user.id:
        raise ApiError(409, "self_archive_denied", "You cannot archive your own account")
    user = await db.get(User, user_id)
    if user is None:
        raise ApiError(404, "member_not_found", "Member not found")
    await ensure_privileged_target_permission(db, principal, user.id)
    await protect_administrator_removal(db, target=user, principal=principal)
    if user.status == "archived":
        return MessageResponse(message="Member is already archived")
    user.status = "archived"
    now = datetime.now(UTC)
    await revoke_member_sessions(db, user.id, now)
    await invalidate_member_access_tokens(db, user.id, now)
    await sync_system_chat_groups(
        db,
        [user],
        created_by_id=principal.user.id,
    )
    record_change(
        db,
        context=event_context(request, principal),
        action="member.archived",
        resource_type="user",
        resource_id=user.id,
        topic="members",
        payload={"user_id": str(user.id), "status": user.status},
    )
    await db.commit()
    return MessageResponse(message="Member archived")


@router.get("/exports/csv")
async def export_members(
    db: DbSession,
    principal: Annotated[Principal, Depends(require_permissions("members.export"))],
) -> StreamingResponse:
    members = (await db.scalars(select(User).order_by(User.surname, User.other_name))).all()
    access = await load_user_access_map(db, [item.id for item in members])
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(
        [
            "staff_id",
            "title",
            "other_name",
            "surname",
            "email",
            "phone_number",
            "academic_rank",
            "status",
            "roles",
        ]
    )
    for member in members:
        writer.writerow(
            [
                csv_safe(member.staff_id),
                csv_safe(member.title),
                csv_safe(member.other_name),
                csv_safe(member.surname),
                csv_safe(member.email),
                csv_safe(member.phone_number),
                csv_safe(member.academic_rank),
                member.status,
                csv_safe(",".join(sorted(access[member.id].roles))),
            ]
        )
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=ug-utag-members.csv"},
    )


@router.get("/import-template.xlsx")
async def member_import_template(
    db: DbSession,
    principal: Annotated[Principal, Depends(require_permissions("members.create"))],
) -> Response:
    del principal
    units = (
        await db.scalars(
            select(OrganizationUnit)
            .where(OrganizationUnit.is_active.is_(True))
            .order_by(OrganizationUnit.name)
        )
    ).all()
    by_id = {unit.id: unit for unit in units}
    departments = sorted(
        (unit for unit in units if unit.unit_type == "department"),
        key=lambda unit: unit.name,
    )

    workbook = Workbook()
    members = workbook.active
    members.title = "Members"
    headers = [
        "staff_id",
        "email",
        "title",
        "other_name",
        "surname",
        "gender",
        "academic_rank",
        "phone_number",
        "college",
        "school",
        "department",
        "roles",
    ]
    members.append(headers)
    members.freeze_panes = "A2"
    members.auto_filter.ref = "A1:L1"
    widths = {
        "A": 18,
        "B": 34,
        "C": 12,
        "D": 22,
        "E": 22,
        "F": 14,
        "G": 26,
        "H": 20,
        "I": 38,
        "J": 42,
        "K": 52,
        "L": 22,
    }
    for column, width in widths.items():
        members.column_dimensions[column].width = width
    for row_number in range(2, 2002):
        members.cell(row=row_number, column=1).number_format = "@"

    header_fill = PatternFill("solid", fgColor="172F4D")
    for cell in members[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center")
    members.row_dimensions[1].height = 24

    instructions = workbook.create_sheet("Instructions")
    instructions.column_dimensions["A"].width = 28
    instructions.column_dimensions["B"].width = 105
    instructions.append(["Requirement", "Guidance"])
    instruction_rows = [
        (
            "Required fields",
            "staff_id, email, other_name, surname, college, school, and department",
        ),
        (
            "Temporary password",
            "The exact staff_id becomes the first password. Members must change "
            "it immediately after sign-in.",
        ),
        (
            "Staff ID format",
            "Keep staff_id cells as text. The template is preformatted to preserve leading zeroes.",
        ),
        (
            "Organization names",
            "Copy exact college, school, and department names from the "
            "Organization Reference sheet.",
        ),
        (
            "Roles",
            "Leave blank for member. Separate multiple roles with commas or semicolons.",
        ),
        (
            "Account and chats",
            "Import activates the account and adds it to UTAG UG plus its "
            "school and department chats.",
        ),
        (
            "Safety",
            "Preview validates every row. Nothing is created until the file "
            "has no reported issues.",
        ),
    ]
    for instruction_row in instruction_rows:
        instructions.append(instruction_row)
    for cell in instructions[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
    for instruction_row in instructions.iter_rows(min_row=2, max_col=2):
        for cell in instruction_row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    reference = workbook.create_sheet("Organization Reference")
    reference.append(["college", "school", "department"])
    reference.freeze_panes = "A2"
    reference.auto_filter.ref = "A1:C1"
    reference.column_dimensions["A"].width = 40
    reference.column_dimensions["B"].width = 46
    reference.column_dimensions["C"].width = 58
    for cell in reference[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
    for department in departments:
        school = by_id.get(department.parent_id) if department.parent_id else None
        college = by_id.get(school.parent_id) if school and school.parent_id else None
        if school and college:
            reference.append([college.name, school.name, department.name])

    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": 'attachment; filename="ug-utag-member-import-template.xlsx"'
        },
    )


@router.get("/organization-reference.csv")
async def organization_reference(
    db: DbSession,
    principal: Annotated[Principal, Depends(require_permissions("members.view"))],
) -> StreamingResponse:
    units = (
        await db.scalars(
            select(OrganizationUnit)
            .where(OrganizationUnit.is_active.is_(True))
            .order_by(OrganizationUnit.unit_type, OrganizationUnit.name)
        )
    ).all()
    by_id = {unit.id: unit for unit in units}
    departments = [unit for unit in units if unit.unit_type == "department"]
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["college", "school", "department"])
    for department in departments:
        school = by_id.get(department.parent_id) if department.parent_id else None
        college = by_id.get(school.parent_id) if school and school.parent_id else None
        writer.writerow(
            [
                csv_safe(college.name if college else ""),
                csv_safe(school.name if school else ""),
                csv_safe(department.name),
            ]
        )
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=ug-organization-reference.csv"},
    )
