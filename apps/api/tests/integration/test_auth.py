import io
from datetime import UTC, datetime, timedelta
from uuid import UUID

from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy import select

from utag_api.database import new_id
from utag_api.models import (
    AccountToken,
    AuditEvent,
    Conversation,
    ConversationMember,
    ExecutiveAppointment,
    MediaAsset,
    OrganizationUnit,
    Role,
    Session,
    User,
    UserPermissionGrant,
    UserRole,
)
from utag_api.routers import auth as auth_router
from utag_api.security import hash_password


async def test_login_me_and_csrf_logout(client: AsyncClient) -> None:
    login = await client.post(
        "/api/v1/auth/login",
        json={"email": "ADMIN@example.edu.gh", "password": "StrongPassword123"},
    )
    assert login.status_code == 200
    payload = login.json()
    assert payload["user"]["full_name"] == "Dr. Ama Mensah"
    assert "settings.manage" in payload["user"]["permissions"]
    assert client.cookies.get("utag_session")
    assert login.headers["cache-control"] == "no-store"
    assert login.headers["x-frame-options"] == "DENY"

    me = await client.get("/api/v1/auth/me")
    assert me.status_code == 200
    assert me.json()["email"] == "admin@example.edu.gh"

    rejected = await client.post("/api/v1/auth/logout")
    assert rejected.status_code == 403
    assert rejected.json()["error"]["code"] == "csrf_failed"

    logout = await client.post(
        "/api/v1/auth/logout", headers={"X-CSRF-Token": payload["csrf_token"]}
    )
    assert logout.status_code == 200


async def test_invalid_credentials_use_generic_message(client: AsyncClient) -> None:
    response = await client.post(
        "/api/v1/auth/login",
        json={"email": "missing@example.edu.gh", "password": "WrongPassword123"},
    )
    assert response.status_code == 401
    assert response.json()["error"]["code"] == "invalid_credentials"


async def test_unknown_account_still_runs_password_verification(
    client: AsyncClient, monkeypatch
) -> None:  # type: ignore[no-untyped-def]
    verified_hashes: list[str] = []

    def fake_verify(password_hash: str, password: str) -> tuple[bool, bool]:
        verified_hashes.append(password_hash)
        return False, False

    monkeypatch.setattr(auth_router, "verify_password", fake_verify)
    response = await client.post(
        "/api/v1/auth/login",
        json={"email": "missing@example.edu.gh", "password": "WrongPassword123"},
    )

    assert response.status_code == 401
    assert verified_hashes == [auth_router.DUMMY_PASSWORD_HASH]


async def test_profile_and_password_workflow(client: AsyncClient) -> None:
    login = await client.post(
        "/api/v1/auth/login",
        json={"email": "admin@example.edu.gh", "password": "StrongPassword123"},
    )
    csrf = login.json()["csrf_token"]
    headers = {"X-CSRF-Token": csrf}

    updated = await client.patch(
        "/api/v1/auth/profile",
        headers=headers,
        json={
            "title": "Prof.",
            "academic_rank": "Associate Professor",
            "phone_number": "+233 20 000 0000",
        },
    )
    assert updated.status_code == 200
    assert updated.json()["full_name"] == "Prof. Ama Mensah"
    assert updated.json()["academic_rank"] == "Associate Professor"

    weak = await client.post(
        "/api/v1/auth/password",
        headers=headers,
        json={"current_password": "StrongPassword123", "new_password": "not-strong-enough"},
    )
    assert weak.status_code == 422
    assert weak.json()["error"]["code"] == "weak_password"

    changed = await client.post(
        "/api/v1/auth/password",
        headers=headers,
        json={"current_password": "StrongPassword123", "new_password": "EvenStrongerPassword456"},
    )
    assert changed.status_code == 200
    assert (await client.get("/api/v1/auth/me")).json()["must_change_password"] is False


async def test_active_executive_can_update_own_public_profile(
    client: AsyncClient, session_factory
) -> None:  # type: ignore[no-untyped-def]
    portrait_id = new_id()
    async with session_factory() as session:
        administrator = await session.scalar(
            select(User).where(User.email == "admin@example.edu.gh")
        )
        assert administrator is not None
        appointment_id = new_id()
        session.add(
            MediaAsset(
                id=portrait_id,
                owner_id=administrator.id,
                storage_key=f"profile/{portrait_id}.png",
                original_filename="president-portrait.png",
                content_type="image/png",
                byte_size=128,
                sha256="a" * 64,
                status="ready",
                is_private=False,
                alt_text="President portrait",
                metadata_json={},
            )
        )
        administrator.profile_media_id = portrait_id
        session.add(
            ExecutiveAppointment(
                id=appointment_id,
                user_id=administrator.id,
                position="President",
                biography_html="<p>Original biography.</p>",
                social_links={},
                is_active=True,
                is_public=True,
            )
        )
        await session.commit()

    public_hidden = await client.get("/api/v1/public/leadership")
    assert public_hidden.status_code == 200
    hidden_profile = next(
        item for item in public_hidden.json() if item["id"] == str(appointment_id)
    )
    assert hidden_profile["show_email"] is False
    assert hidden_profile["email"] is None

    login = await client.post(
        "/api/v1/auth/login",
        json={"email": "admin@example.edu.gh", "password": "StrongPassword123"},
    )
    assert login.status_code == 200
    assert login.json()["user"]["must_complete_executive_profile"] is True
    headers = {"X-CSRF-Token": login.json()["csrf_token"]}

    blocked_chat = await client.get("/api/v1/chat/conversations")
    assert blocked_chat.status_code == 403
    assert blocked_chat.json()["error"]["code"] == "executive_profile_required"

    current = await client.get("/api/v1/auth/executive-profile")
    assert current.status_code == 200
    assert current.json()["position"] == "President"

    too_short = await client.patch(
        "/api/v1/auth/executive-profile",
        headers=headers,
        json={"biography_html": "<p>Hi</p>", "social_links": {}},
    )
    assert too_short.status_code == 422
    assert too_short.json()["error"]["code"] == "executive_biography_required"

    invalid_link = await client.patch(
        "/api/v1/auth/executive-profile",
        headers=headers,
        json={
            "biography_html": "<p>Biography for the public leadership page.</p>",
            "social_links": {"linkedin": "javascript:alert('unsafe')"},
        },
    )
    assert invalid_link.status_code == 422

    updated = await client.patch(
        "/api/v1/auth/executive-profile",
        headers=headers,
        json={
            "portfolio": "Member welfare and conditions of service",
            "summary": "Working for fair conditions across campus.",
            "biography_html": (
                "<p>Serving <strong>UG UTAG members</strong>.</p><script>alert('unsafe')</script>"
            ),
            "social_links": {
                "linkedin": "https://www.linkedin.com/in/ug-utag-test",
                "facebook": "",
            },
            "show_email": True,
            "show_phone": False,
        },
    )
    assert updated.status_code == 200
    assert updated.json()["biography_html"] == ("<p>Serving <strong>UG UTAG members</strong>.</p>")
    assert updated.json()["social_links"] == {
        "linkedin": "https://www.linkedin.com/in/ug-utag-test"
    }
    assert updated.json()["portfolio"] == "Member welfare and conditions of service"
    assert updated.json()["summary"] == "Working for fair conditions across campus."
    assert updated.json()["show_email"] is True
    assert updated.json()["show_phone"] is False

    me = await client.get("/api/v1/auth/me")
    assert me.status_code == 200
    assert me.json()["must_complete_executive_profile"] is False

    conversations = await client.get("/api/v1/chat/conversations")
    assert conversations.status_code == 200

    public = await client.get("/api/v1/public/leadership")
    assert public.status_code == 200
    profile = next(item for item in public.json() if item["id"] == str(appointment_id))
    assert profile["biography_html"] == updated.json()["biography_html"]
    assert profile["show_email"] is True
    assert profile["email"] == "admin@example.edu.gh"
    assert profile["phone_number"] is None

    withheld = await client.patch(
        "/api/v1/auth/executive-profile",
        headers=headers,
        json={
            "portfolio": updated.json()["portfolio"],
            "summary": updated.json()["summary"],
            "biography_html": updated.json()["biography_html"],
            "social_links": updated.json()["social_links"],
            "show_email": False,
            "show_phone": False,
        },
    )
    assert withheld.status_code == 200
    public_withheld = await client.get("/api/v1/public/leadership")
    withheld_profile = next(
        item for item in public_withheld.json() if item["id"] == str(appointment_id)
    )
    assert withheld_profile["show_email"] is False
    assert withheld_profile["email"] is None

    async with session_factory() as session:
        event = await session.scalar(
            select(AuditEvent)
            .where(
                AuditEvent.action == "executive.profile_updated",
                AuditEvent.resource_id == appointment_id,
            )
            .order_by(AuditEvent.created_at.desc())
        )
        assert event is not None


async def test_administrator_can_correct_a_member_email_safely(
    client: AsyncClient,
    session_factory,
) -> None:  # type: ignore[no-untyped-def]
    login = await client.post(
        "/api/v1/auth/login",
        json={"email": "admin@example.edu.gh", "password": "StrongPassword123"},
    )
    headers = {"X-CSRF-Token": login.json()["csrf_token"]}
    member_id = new_id()
    now = datetime.now(UTC)

    async with session_factory() as session:
        session.add_all(
            [
                User(
                    id=member_id,
                    email="incorrect.member@example.edu.gh",
                    password_hash=hash_password("MemberPassword123"),
                    status="active",
                    email_verified=True,
                    other_name="Kwesi",
                    surname="Arthur",
                ),
                Session(
                    id=new_id(),
                    user_id=member_id,
                    token_hash="a" * 64,
                    csrf_hash="b" * 64,
                    user_agent="member-browser",
                    ip_prefix="127.0.0",
                    created_at=now,
                    last_seen_at=now,
                    expires_at=now + timedelta(days=1),
                ),
                AccountToken(
                    id=new_id(),
                    user_id=member_id,
                    kind="password_reset",
                    token_hash="c" * 64,
                    created_at=now,
                    expires_at=now + timedelta(minutes=30),
                ),
            ]
        )
        await session.commit()

    updated = await client.patch(
        f"/api/v1/members/{member_id}",
        headers=headers,
        json={"email": " Corrected.Member@Example.EDU.GH "},
    )
    assert updated.status_code == 200
    assert updated.json()["email"] == "corrected.member@example.edu.gh"
    assert updated.json()["email_verified"] is False

    async with session_factory() as session:
        revoked_session = await session.scalar(
            select(Session).where(Session.token_hash == "a" * 64)
        )
        invalidated_token = await session.scalar(
            select(AccountToken).where(AccountToken.token_hash == "c" * 64)
        )
        audit_event = await session.scalar(
            select(AuditEvent)
            .where(
                AuditEvent.resource_id == member_id,
                AuditEvent.action == "member.updated",
            )
            .order_by(AuditEvent.created_at.desc())
        )
        assert revoked_session is not None and revoked_session.revoked_at is not None
        assert invalidated_token is not None and invalidated_token.used_at is not None
        assert audit_event is not None
        assert audit_event.changes["email"]["to"] == "corrected.member@example.edu.gh"

        unchanged_email_session = Session(
            id=new_id(),
            user_id=member_id,
            token_hash="d" * 64,
            csrf_hash="e" * 64,
            user_agent="member-browser",
            ip_prefix="127.0.0",
            created_at=now,
            last_seen_at=now,
            expires_at=now + timedelta(days=1),
        )
        session.add(unchanged_email_session)
        await session.commit()

    unchanged = await client.patch(
        f"/api/v1/members/{member_id}",
        headers=headers,
        json={
            "email": "corrected.member@example.edu.gh",
            "phone_number": "+233200000001",
        },
    )
    assert unchanged.status_code == 200
    async with session_factory() as session:
        unchanged_email_session = await session.scalar(
            select(Session).where(Session.token_hash == "d" * 64)
        )
        assert unchanged_email_session is not None
        assert unchanged_email_session.revoked_at is None

    duplicate = await client.patch(
        f"/api/v1/members/{member_id}",
        headers=headers,
        json={"email": "ADMIN@EXAMPLE.EDU.GH"},
    )
    assert duplicate.status_code == 409
    assert duplicate.json()["error"]["code"] == "email_exists"

    invalid = await client.patch(
        f"/api/v1/members/{member_id}",
        headers=headers,
        json={"email": "not-an-email"},
    )
    assert invalid.status_code == 422

    old_email_login = await client.post(
        "/api/v1/auth/login",
        json={
            "email": "incorrect.member@example.edu.gh",
            "password": "MemberPassword123",
        },
    )
    assert old_email_login.status_code == 401
    new_email_login = await client.post(
        "/api/v1/auth/login",
        json={
            "email": "corrected.member@example.edu.gh",
            "password": "MemberPassword123",
        },
    )
    assert new_email_login.status_code == 200


async def test_member_management_archives_without_deleting(
    client: AsyncClient, monkeypatch, session_factory
) -> None:  # type: ignore[no-untyped-def]
    login = await client.post(
        "/api/v1/auth/login",
        json={"email": "admin@example.edu.gh", "password": "StrongPassword123"},
    )
    payload = login.json()
    headers = {"X-CSRF-Token": payload["csrf_token"]}

    created = await client.post(
        "/api/v1/members",
        headers=headers,
        json={
            "email": "new.member@example.edu.gh",
            "title": "Dr.",
            "other_name": "Kwame",
            "surname": "Asante",
            "roles": ["member"],
            "send_invitation": False,
        },
    )
    assert created.status_code == 201
    member_id = created.json()["id"]

    profile_asset_id = new_id()
    invalid_profile_asset_id = new_id()
    async with session_factory() as session:
        administrator = await session.get(User, UUID(payload["user"]["id"]))
        assert administrator is not None
        session.add_all(
            [
                MediaAsset(
                    id=profile_asset_id,
                    owner_id=administrator.id,
                    storage_key=f"profile/{profile_asset_id}.png",
                    original_filename="executive-portrait.png",
                    content_type="image/png",
                    byte_size=128,
                    sha256="2" * 64,
                    status="ready",
                    is_private=False,
                    alt_text="Executive portrait",
                    metadata_json={},
                ),
                MediaAsset(
                    id=invalid_profile_asset_id,
                    owner_id=administrator.id,
                    storage_key=f"profile/{invalid_profile_asset_id}.pdf",
                    original_filename="not-a-portrait.pdf",
                    content_type="application/pdf",
                    byte_size=128,
                    sha256="3" * 64,
                    status="ready",
                    is_private=False,
                    alt_text=None,
                    metadata_json={},
                ),
            ],
        )
        await session.commit()

    updated = await client.patch(
        f"/api/v1/members/{member_id}",
        headers=headers,
        json={
            "academic_rank": "Senior Lecturer",
            "roles": ["member", "editor"],
            "profile_media_id": str(profile_asset_id),
        },
    )
    assert updated.status_code == 200
    assert updated.json()["roles"] == ["editor", "member"]
    assert updated.json()["profile_media_id"] == str(profile_asset_id)

    invalid_profile = await client.patch(
        f"/api/v1/members/{member_id}",
        headers=headers,
        json={"profile_media_id": str(invalid_profile_asset_id)},
    )
    assert invalid_profile.status_code == 422
    assert invalid_profile.json()["error"]["code"] == "profile_image_invalid"

    status_bypass = await client.patch(
        f"/api/v1/members/{member_id}",
        headers=headers,
        json={"status": "suspended"},
    )
    assert status_bypass.status_code == 409
    assert status_bypass.json()["error"]["code"] == "member_lifecycle_action_required"

    async with session_factory() as session:
        member = await session.get(User, UUID(member_id))
        assert member is not None
        member.email_verified = True
        await session.commit()

    monkeypatch.setattr("utag_api.worker.tasks.send_password_reset.delay", lambda *_: None)
    reset = await client.post(f"/api/v1/members/{member_id}/password-reset", headers=headers)
    assert reset.status_code == 200
    refreshed = await client.get(f"/api/v1/members/{member_id}")
    assert refreshed.json()["must_change_password"] is True
    async with session_factory() as session:
        reset_tokens = (
            await session.scalars(
                select(AccountToken).where(
                    AccountToken.user_id == UUID(member_id),
                    AccountToken.kind == "password_reset",
                    AccountToken.used_at.is_(None),
                )
            )
        ).all()
        assert len(reset_tokens) == 1

    deactivated = await client.post(
        f"/api/v1/members/{member_id}/deactivate",
        headers=headers,
        json={"reason": "Member is temporarily away from the association"},
    )
    assert deactivated.status_code == 200
    assert (await client.get(f"/api/v1/members/{member_id}")).json()["status"] == "suspended"
    access_while_suspended = await client.post(
        f"/api/v1/members/{member_id}/access-link", headers=headers
    )
    assert access_while_suspended.status_code == 409

    reactivated = await client.post(f"/api/v1/members/{member_id}/reactivate", headers=headers)
    assert reactivated.status_code == 200
    assert (await client.get(f"/api/v1/members/{member_id}")).json()["status"] == "active"

    exported = await client.get("/api/v1/members/exports/csv")
    assert exported.status_code == 200
    assert "new.member@example.edu.gh" in exported.text
    reference = await client.get("/api/v1/members/organization-reference.csv")
    assert reference.status_code == 200
    assert reference.text.startswith("college,school,department")

    archived = await client.delete(f"/api/v1/members/{member_id}", headers=headers)
    assert archived.status_code == 200
    retained = await client.get(f"/api/v1/members/{member_id}")
    assert retained.status_code == 200
    assert retained.json()["status"] == "archived"

    self_archive = await client.delete(f"/api/v1/members/{payload['user']['id']}", headers=headers)
    assert self_archive.status_code == 409
    assert self_archive.json()["error"]["code"] == "self_archive_denied"


async def test_member_and_executive_creation_accept_inline_profile_images(
    client: AsyncClient, session_factory
) -> None:  # type: ignore[no-untyped-def]
    login = await client.post(
        "/api/v1/auth/login",
        json={"email": "admin@example.edu.gh", "password": "StrongPassword123"},
    )
    payload = login.json()
    headers = {"X-CSRF-Token": payload["csrf_token"]}
    public_image_id = new_id()
    private_image_id = new_id()
    async with session_factory() as session:
        administrator = await session.get(User, UUID(payload["user"]["id"]))
        assert administrator is not None
        session.add_all(
            [
                MediaAsset(
                    id=public_image_id,
                    owner_id=administrator.id,
                    storage_key=f"profile/{public_image_id}.png",
                    original_filename="new-executive.png",
                    content_type="image/png",
                    byte_size=256,
                    sha256="4" * 64,
                    status="ready",
                    is_private=False,
                    alt_text="New executive portrait",
                    metadata_json={},
                ),
                MediaAsset(
                    id=private_image_id,
                    owner_id=administrator.id,
                    storage_key=f"profile/{private_image_id}.png",
                    original_filename="private.png",
                    content_type="image/png",
                    byte_size=256,
                    sha256="5" * 64,
                    status="ready",
                    is_private=True,
                    alt_text="Private portrait",
                    metadata_json={},
                ),
            ]
        )
        await session.commit()

    media_status = await client.get(f"/api/v1/media/{public_image_id}")
    assert media_status.status_code == 200
    assert media_status.json()["status"] == "ready"

    created_member = await client.post(
        "/api/v1/members",
        headers=headers,
        json={
            "email": "inline.photo@example.edu.gh",
            "other_name": "Inline",
            "surname": "Photo",
            "profile_media_id": str(public_image_id),
            "roles": ["member"],
            "send_invitation": False,
        },
    )
    assert created_member.status_code == 201
    assert created_member.json()["profile_media_id"] == str(public_image_id)

    missing_photo = await client.post(
        "/api/v1/executives",
        headers=headers,
        json={
            "user_id": created_member.json()["id"],
            "position": "Executive Member",
        },
    )
    assert missing_photo.status_code == 422

    created_executive = await client.post(
        "/api/v1/executives",
        headers=headers,
        json={
            "user_id": created_member.json()["id"],
            "profile_media_id": str(public_image_id),
            "position": "Executive Member",
        },
    )
    assert created_executive.status_code == 201
    assert created_executive.json()["profile_media_id"] == str(public_image_id)

    rejected_member = await client.post(
        "/api/v1/members",
        headers=headers,
        json={
            "email": "private.photo@example.edu.gh",
            "other_name": "Private",
            "surname": "Photo",
            "profile_media_id": str(private_image_id),
            "roles": ["member"],
            "send_invitation": False,
        },
    )
    assert rejected_member.status_code == 422
    assert rejected_member.json()["error"]["code"] == "profile_image_invalid"


async def test_individual_permission_grants_merge_with_role_access(
    client: AsyncClient, session_factory
) -> None:  # type: ignore[no-untyped-def]
    admin_login = await client.post(
        "/api/v1/auth/login",
        json={"email": "admin@example.edu.gh", "password": "StrongPassword123"},
    )
    admin_headers = {"X-CSRF-Token": admin_login.json()["csrf_token"]}

    catalog = await client.get("/api/v1/members/permission-options")
    assert catalog.status_code == 200
    option_keys = {item["key"] for item in catalog.json()}
    assert {"content.edit", "media.manage"} <= option_keys
    assert "members.roles" not in option_keys
    assert "members.permissions" not in option_keys
    assert "records.delete" not in option_keys

    created = await client.post(
        "/api/v1/members",
        headers=admin_headers,
        json={
            "email": "custom.access@example.edu.gh",
            "other_name": "Adwoa",
            "surname": "Owusu",
            "roles": ["member"],
            "send_invitation": False,
        },
    )
    assert created.status_code == 201
    member_id = created.json()["id"]

    non_delegable = await client.put(
        f"/api/v1/members/{member_id}/permissions",
        headers=admin_headers,
        json={"permissions": ["members.roles"]},
    )
    assert non_delegable.status_code == 422
    assert non_delegable.json()["error"]["code"] == "permission_not_delegable"

    granted = await client.put(
        f"/api/v1/members/{member_id}/permissions",
        headers=admin_headers,
        json={"permissions": ["media.manage", "content.edit", "content.edit"]},
    )
    assert granted.status_code == 200
    grant_payload = granted.json()
    assert grant_payload["roles"] == ["member"]
    assert grant_payload["extra_permissions"] == ["content.edit", "media.manage"]
    assert {"chat.use", "dashboard.view", "documents.view"} <= set(
        grant_payload["role_permissions"]
    )
    assert set(grant_payload["effective_permissions"]) == set(grant_payload["role_permissions"]) | {
        "content.edit",
        "media.manage",
    }

    async with session_factory() as session:
        member = await session.get(User, UUID(member_id))
        assert member is not None
        member.password_hash = hash_password("CustomAccessPassword123")
        member.email_verified = True
        member.must_change_password = False
        await session.commit()

    member_login = await client.post(
        "/api/v1/auth/login",
        json={
            "email": "custom.access@example.edu.gh",
            "password": "CustomAccessPassword123",
        },
    )
    assert member_login.status_code == 200
    assert {"content.edit", "media.manage"} <= set(member_login.json()["user"]["permissions"])
    member_headers = {"X-CSRF-Token": member_login.json()["csrf_token"]}
    denied = await client.put(
        f"/api/v1/members/{member_id}/permissions",
        headers=member_headers,
        json={"permissions": []},
    )
    assert denied.status_code == 403
    assert denied.json()["error"]["code"] == "permission_denied"

    admin_login = await client.post(
        "/api/v1/auth/login",
        json={"email": "admin@example.edu.gh", "password": "StrongPassword123"},
    )
    cleared = await client.put(
        f"/api/v1/members/{member_id}/permissions",
        headers={"X-CSRF-Token": admin_login.json()["csrf_token"]},
        json={"permissions": []},
    )
    assert cleared.status_code == 200
    assert cleared.json()["extra_permissions"] == []
    assert cleared.json()["effective_permissions"] == cleared.json()["role_permissions"]

    async with session_factory() as session:
        remaining_grants = (
            await session.scalars(
                select(UserPermissionGrant).where(UserPermissionGrant.user_id == UUID(member_id))
            )
        ).all()
        audit_events = (
            await session.scalars(
                select(AuditEvent).where(
                    AuditEvent.resource_id == UUID(member_id),
                    AuditEvent.action == "member.permissions.updated",
                )
            )
        ).all()
        assert remaining_grants == []
        assert len(audit_events) == 2


async def test_secretary_cannot_assign_roles_or_manage_administrators(
    client: AsyncClient, session_factory
) -> None:  # type: ignore[no-untyped-def]
    async with session_factory() as session:
        secretary_role = await session.scalar(select(Role).where(Role.key == "secretary"))
        administrator = await session.scalar(
            select(User).where(User.email == "admin@example.edu.gh")
        )
        assert secretary_role is not None
        assert administrator is not None
        secretary = User(
            id=new_id(),
            email="secretary@example.edu.gh",
            password_hash=hash_password("SecretaryPassword123"),
            status="active",
            email_verified=True,
            other_name="Akosua",
            surname="Boateng",
        )
        session.add(secretary)
        await session.flush()
        session.add(
            UserRole(
                id=new_id(),
                user_id=secretary.id,
                role_id=secretary_role.id,
                assigned_by_id=administrator.id,
                assigned_at=datetime.now(UTC),
            )
        )
        administrator_id = administrator.id
        await session.commit()

    login = await client.post(
        "/api/v1/auth/login",
        json={
            "email": "secretary@example.edu.gh",
            "password": "SecretaryPassword123",
        },
    )
    assert login.status_code == 200
    permissions = login.json()["user"]["permissions"]
    assert "members.create" in permissions
    assert "members.credentials" in permissions
    assert "members.roles" not in permissions
    assert "members.permissions" not in permissions
    assert "records.delete" not in permissions
    headers = {"X-CSRF-Token": login.json()["csrf_token"]}

    created = await client.post(
        "/api/v1/members",
        headers=headers,
        json={
            "email": "ordinary.member@example.edu.gh",
            "other_name": "Kojo",
            "surname": "Addo",
            "roles": ["member"],
            "send_invitation": False,
        },
    )
    assert created.status_code == 201

    escalated = await client.post(
        "/api/v1/members",
        headers=headers,
        json={
            "email": "unauthorized.admin@example.edu.gh",
            "other_name": "Bad",
            "surname": "Grant",
            "roles": ["administrator"],
            "send_invitation": False,
        },
    )
    assert escalated.status_code == 403
    assert escalated.json()["error"]["code"] == "role_assignment_denied"

    administrator_update = await client.patch(
        f"/api/v1/members/{administrator_id}",
        headers=headers,
        json={"phone_number": "+233200000001"},
    )
    assert administrator_update.status_code == 403
    assert administrator_update.json()["error"]["code"] == "administrator_management_denied"


async def test_member_import_creates_staff_password_accounts_and_chat_groups(
    client: AsyncClient,
    session_factory,
) -> None:  # type: ignore[no-untyped-def]
    async with session_factory() as session:
        college = OrganizationUnit(
            id=new_id(),
            unit_type="college",
            name="College of Import Studies",
            slug="college-import-studies",
            is_active=True,
        )
        school = OrganizationUnit(
            id=new_id(),
            unit_type="school",
            name="School of Import Studies",
            slug="school-import-studies",
            parent_id=college.id,
            is_active=True,
        )
        department = OrganizationUnit(
            id=new_id(),
            unit_type="department",
            name="Department of Import Studies",
            slug="department-import-studies",
            parent_id=school.id,
            is_active=True,
        )
        session.add_all([college, school, department])
        await session.commit()

    login = await client.post(
        "/api/v1/auth/login",
        json={"email": "admin@example.edu.gh", "password": "StrongPassword123"},
    )
    headers = {"X-CSRF-Token": login.json()["csrf_token"]}
    file_content = (
        b"email,staff_id,other_name,surname,academic_rank,college,school,department,roles\n"
        b"bulk.member@example.edu.gh,UG1,Efua,Owusu,Lecturer,"
        b"College of Import Studies,School of Import Studies,"
        b"Department of Import Studies,member\n"
    )

    missing_staff_id = await client.post(
        "/api/v1/members/import",
        headers=headers,
        files={
            "file": (
                "members.csv",
                file_content.replace(b"UG1", b""),
                "text/csv",
            )
        },
        data={"dry_run": "true"},
    )
    assert missing_staff_id.status_code == 200
    assert missing_staff_id.json()["valid_rows"] == 0
    assert missing_staff_id.json()["issues"][0]["field"] == "staff_id"

    existing_email = await client.post(
        "/api/v1/members/import",
        headers=headers,
        files={
            "file": (
                "members.csv",
                file_content.replace(
                    b"bulk.member@example.edu.gh",
                    b"ADMIN@example.edu.gh",
                ),
                "text/csv",
            )
        },
        data={"dry_run": "true"},
    )
    assert existing_email.status_code == 200
    assert existing_email.json()["valid_rows"] == 0
    assert existing_email.json()["issues"][0]["field"] == "email"

    preview = await client.post(
        "/api/v1/members/import",
        headers=headers,
        data={"dry_run": "true"},
        files={"file": ("members.csv", file_content, "text/csv")},
    )
    assert preview.status_code == 200
    assert preview.json()["valid_rows"] == 1
    assert preview.json()["imported_rows"] == 0
    assert preview.json()["issues"] == []
    assert preview.json()["preview"][0]["chat_groups"] == [
        "UTAG UG",
        "School: School of Import Studies",
        "Department: Department of Import Studies",
    ]

    committed = await client.post(
        "/api/v1/members/import",
        headers=headers,
        data={"dry_run": "false"},
        files={"file": ("members.csv", file_content, "text/csv")},
    )
    assert committed.status_code == 200
    assert committed.json()["imported_rows"] == 1
    assert committed.json()["chat_groups_created"] == 3
    assert committed.json()["chat_memberships_added"] == 3

    members = await client.get("/api/v1/members", params={"q": "bulk.member"})
    assert members.status_code == 200
    imported = members.json()["items"][0]
    assert imported["status"] == "active"
    assert imported["must_change_password"] is True
    assert imported["roles"] == ["member"]

    first_login = await client.post(
        "/api/v1/auth/login",
        json={"email": "bulk.member@example.edu.gh", "password": "UG1"},
    )
    assert first_login.status_code == 200
    assert first_login.json()["user"]["must_change_password"] is True
    member_headers = {"X-CSRF-Token": first_login.json()["csrf_token"]}

    blocked_chat = await client.get("/api/v1/chat/conversations")
    assert blocked_chat.status_code == 403
    assert blocked_chat.json()["error"]["code"] == "password_change_required"

    changed = await client.post(
        "/api/v1/auth/password",
        headers=member_headers,
        json={
            "current_password": "UG1",
            "new_password": "PrivatePassword456",
        },
    )
    assert changed.status_code == 200

    conversations = await client.get("/api/v1/chat/conversations")
    assert conversations.status_code == 200
    assert {item["title"] for item in conversations.json()} == {
        "UTAG UG",
        "School: School of Import Studies",
        "Department: Department of Import Studies",
    }
    utag_group = next(item for item in conversations.json() if item["title"] == "UTAG UG")
    managed_leave = await client.delete(
        f"/api/v1/chat/conversations/{utag_group['id']}/members/{imported['id']}",
        headers=member_headers,
    )
    assert managed_leave.status_code == 409
    assert managed_leave.json()["error"]["code"] == "system_chat_membership_managed"

    admin_login = await client.post(
        "/api/v1/auth/login",
        json={"email": "admin@example.edu.gh", "password": "StrongPassword123"},
    )
    admin_headers = {"X-CSRF-Token": admin_login.json()["csrf_token"]}
    deactivated = await client.post(
        f"/api/v1/members/{imported['id']}/deactivate",
        headers=admin_headers,
        json={"reason": "Membership paused for integration testing"},
    )
    assert deactivated.status_code == 200
    async with session_factory() as session:
        managed_memberships = (
            await session.scalars(
                select(ConversationMember)
                .join(
                    Conversation,
                    Conversation.id == ConversationMember.conversation_id,
                )
                .where(
                    ConversationMember.user_id == UUID(imported["id"]),
                    Conversation.direct_key.like("system:%"),
                )
            )
        ).all()
        assert len(managed_memberships) == 3
        assert all(item.left_at is not None for item in managed_memberships)

    reactivated = await client.post(
        f"/api/v1/members/{imported['id']}/reactivate",
        headers=admin_headers,
    )
    assert reactivated.status_code == 200
    async with session_factory() as session:
        active_memberships = (
            await session.scalars(
                select(ConversationMember)
                .join(
                    Conversation,
                    Conversation.id == ConversationMember.conversation_id,
                )
                .where(
                    ConversationMember.user_id == UUID(imported["id"]),
                    Conversation.direct_key.like("system:%"),
                    ConversationMember.left_at.is_(None),
                )
            )
        ).all()
        assert len(active_memberships) == 3


async def test_member_import_template_is_a_valid_workbook(
    client: AsyncClient,
) -> None:
    login = await client.post(
        "/api/v1/auth/login",
        json={"email": "admin@example.edu.gh", "password": "StrongPassword123"},
    )
    assert login.status_code == 200

    response = await client.get("/api/v1/members/import-template.xlsx")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "ug-utag-member-import-template.xlsx" in response.headers["content-disposition"]

    workbook = load_workbook(io.BytesIO(response.content), read_only=True)
    try:
        assert workbook.sheetnames == [
            "Members",
            "Instructions",
            "Organization Reference",
        ]
        assert list(next(workbook["Members"].iter_rows(values_only=True))) == [
            "staff_id",
            "email",
            "title",
            "other_name",
            "surname",
            "gender",
            "academic_rank",
            "phone_number",
            "college",
            "school",
            "department",
            "roles",
        ]
    finally:
        workbook.close()


async def test_carousel_archive_and_permanent_delete_both_retain_media(
    client: AsyncClient, session_factory
) -> None:  # type: ignore[no-untyped-def]
    login = await client.post(
        "/api/v1/auth/login",
        json={"email": "admin@example.edu.gh", "password": "StrongPassword123"},
    )
    headers = {"X-CSRF-Token": login.json()["csrf_token"]}
    asset_id = new_id()
    async with session_factory() as session:
        administrator = await session.scalar(
            select(User).where(User.email == "admin@example.edu.gh")
        )
        assert administrator is not None
        session.add(
            MediaAsset(
                id=asset_id,
                owner_id=administrator.id,
                storage_key=f"carousel/{asset_id}.png",
                original_filename="homepage.png",
                content_type="image/png",
                byte_size=128,
                sha256="0" * 64,
                status="ready",
                is_private=False,
                alt_text="UG UTAG members",
                metadata_json={},
            )
        )
        await session.commit()

    media_search = await client.get(
        "/api/v1/media", params={"status": "ready", "q": "UTAG members"}
    )
    assert media_search.status_code == 200
    assert media_search.json()["items"][0]["id"] == str(asset_id)
    no_media_match = await client.get("/api/v1/media", params={"q": "not-present"})
    assert no_media_match.status_code == 200
    assert no_media_match.json()["items"] == []

    created = await client.post(
        "/api/v1/admin/carousel",
        headers=headers,
        json={
            "title": "Member solidarity",
            "description": (
                "<p>A <strong>public homepage</strong> highlight."
                "<script>alert('unsafe')</script></p>"
            ),
            "media_asset_id": str(asset_id),
            "order": 1,
            "is_published": True,
        },
    )
    assert created.status_code == 201
    slide_id = created.json()["id"]
    assert "<strong>public homepage</strong>" in created.json()["description"]
    assert "script" not in created.json()["description"]

    listing = await client.get("/api/v1/admin/carousel")
    assert listing.status_code == 200
    assert listing.json()[0]["title"] == "Member solidarity"

    public_settings = await client.get("/api/v1/public/settings")
    assert public_settings.status_code == 200
    assert public_settings.json()["site.carousel"]["slides"][0]["id"] == slide_id

    updated = await client.patch(
        f"/api/v1/admin/carousel/{slide_id}",
        headers=headers,
        json={
            "title": "Updated member solidarity",
            "description": '<p onclick="bad()">Updated message.</p>',
        },
    )
    assert updated.status_code == 200
    assert updated.json()["title"] == "Updated member solidarity"
    assert updated.json()["description"] == "<p>Updated message.</p>"

    archived = await client.delete(f"/api/v1/admin/carousel/{slide_id}", headers=headers)
    assert archived.status_code == 200
    assert (await client.get("/api/v1/public/settings")).json()["site.carousel"]["slides"] == []
    retained = (await client.get("/api/v1/admin/carousel")).json()[0]
    assert retained["archived"] is True
    assert retained["is_published"] is False

    deleted = await client.delete(
        f"/api/v1/admin/carousel/{slide_id}/permanent",
        headers=headers,
    )
    assert deleted.status_code == 200
    assert (await client.get("/api/v1/admin/carousel")).json() == []
    async with session_factory() as session:
        assert await session.get(MediaAsset, asset_id) is not None
